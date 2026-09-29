"""Exporta la tabla de catalogación: columna A = registro completo; columnas siguientes = etiquetas del
catálogo (vacías); al final, la ubicación exacta en la fuente para corregir a mano.

Lee las tablas registros / registros_fuente / registros_exclusion que escribe extraer_registros.py.
Uso: python3 tools/exportar_catalogo.py --xlsx salida.xlsx [--db data/rules.sqlite]
"""

from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

APD = Path(__file__).resolve().parents[1]

TAXONOMIA = APD / "TAXONOMIA.md"


def leer_taxonomia(ruta=TAXONOMIA):
    """Facetas aprobadas (decisión 10): [("N. Nombre", "valor · valor · ...")], leídas de TAXONOMIA.md."""
    import re
    t = ruta.read_text(encoding="utf-8")
    out = []
    for b in re.split(r"\n### ", t)[1:]:
        b = b.split("\n## ")[0]
        cab = b.split("\n", 1)[0]
        nombre = cab.split(" — ")[0].split(" (")[0].strip()
        filas = [l for l in b.split("\n") if l.startswith("|") and not l.startswith("|---")]
        hdr = [c.strip() for c in filas[0].strip("|").split("|")] if filas else []
        vals = []
        for f in filas[1:]:
            cel = [c.strip() for c in f.strip("|").split("|")]
            if hdr[:2] == ["Rama", "Valores"] or hdr[0] == "Etapa":
                vals += [x.strip() for x in cel[1].split("·")]
            elif hdr[0] == "Rama":
                vals.append(cel[1])
            else:
                vals.append(cel[0])
        out.append((nombre, " · ".join(dict.fromkeys(vals))))
    return out


CATALOGO = leer_taxonomia()

UBICACION = ["Documento", "Página / sección", "Párrafo nº (de N)", "Líneas", "ID registro"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default=str(APD / "data/rules.sqlite"))
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sin-duplicados", action="store_true",
                    help="quita los duplicados de dd_eliminado (depurar_duplicados.py); cada conservado recibe las ubicaciones de los que contiene")
    a = ap.parse_args()
    con = sqlite3.connect(a.db)
    regs = con.execute("select id, orden, texto from registros order by orden").fetchall()
    contexto = dict(con.execute("select id, contexto from registros"))
    fuentes: dict[str, list] = {}
    for rid, arch, li, lf, sec, par, tot in con.execute(
            "select registro_id, archivo, linea_ini, linea_fin, seccion, parrafo, parrafos_en_documento from registros_fuente"):
        fuentes.setdefault(rid, []).append((arch, li, lf, sec, par, tot))
    eliminados = []
    if a.sin_duplicados:
        if not (APD / "data/depuracion/final.json").exists():
            raise SystemExit("falta el cierre: correr 'depurar_duplicados.py final' antes de exportar")
        texto_de = {rid: t for rid, _, t in regs}
        for rid, cont, g, raz in con.execute("select id, contenedor, grupo, razon from dd_eliminado order by grupo"):
            eliminados.append((g, rid, cont, raz))
            # la ubicación del duplicado eliminado pasa al original conservado que lo contiene
            fuentes[cont] = fuentes.get(cont, []) + fuentes.get(rid, [])
        fuera = {rid for _, rid, _, _ in eliminados}
        regs = [r for r in regs if r[0] not in fuera]
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    sys.path.insert(0, str(APD / "tools"))
    import depurar_duplicados as dd
    sk, rp, ht = dd.verdad_del_usuario()
    etiqueta = {"ok": "sí, literal", "mal": "DISTINTO a lo que subiste", None: "sin copia en lo que subiste"}
    def _resp(rid, t):
        docs = [f[0] for f in fuentes.get(rid, [])]
        if docs and all(d.startswith("rules.sqlite:") for d in docs):
            import sqlite3 as _s
            z = _s.connect(rp_db)
            hay = z.execute("select 1 from reglas where texto = ? limit 1", (t,)).fetchone()
            return "sí, literal en rules.sqlite de tu zip" if hay else "DISTINTO a lo que subiste"
        return etiqueta[dd.estado_respaldo(con, t, rid, sk, rp, ht)]
    import tempfile, zipfile as _z, glob as _g
    _tmp = tempfile.mkdtemp()
    _z.ZipFile(_g.glob("/root/.claude/uploads/*/*AI-Production-Director-completo.zip")[0]).extract("rules.sqlite", _tmp)
    rp_db = str(Path(_tmp) / "rules.sqlite")
    respaldo = {rid: _resp(rid, t) for rid, _, t in regs}
    cab = ["Registro (copia literal de la fuente)"] + [c for c, _ in CATALOGO] + UBICACION + \
        ["Contexto literal (encabezado de tabla o de bloque)", "Respaldo en tu zip"]
    ws.append(cab)
    for rid, orden, texto in regs:
        fs = fuentes.get(rid, [])
        doc = " | ".join(f[0] for f in fs)
        sec = " | ".join(f[3] or "" for f in fs)
        par = " | ".join(f"{f[4]} de {f[5]}" for f in fs)
        lin = " | ".join((f"{f[1]}–{f[2]}" if f[1] and f[2] != f[1] else str(f[1])) if f[1] else "objeto JSON/Python" for f in fs)
        ws.append([texto] + [""] * len(CATALOGO) + [doc, sec, par, lin, rid, contexto.get(rid) or "", respaldo[rid]])
    bold = Font(bold=True)
    fill_cat = PatternFill("solid", fgColor="FFF2CC")
    fill_ubi = PatternFill("solid", fgColor="DDEBF7")
    for col in range(1, len(cab) + 1):
        c = ws.cell(row=1, column=col)
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if 2 <= col <= 1 + len(CATALOGO):
            c.fill = fill_cat
        elif col > 1 + len(CATALOGO):
            c.fill = fill_ubi
    ws.column_dimensions["A"].width = 110
    for col in range(2, 2 + len(CATALOGO)):
        ws.column_dimensions[get_column_letter(col)].width = 16
    for i, w in enumerate([48, 48, 14, 12, 14]):
        ws.column_dimensions[get_column_letter(2 + len(CATALOGO) + i)].width = w
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{ws.max_row}"

    if eliminados:
        wd = wb.create_sheet("Duplicados eliminados")
        wd.append(["Grupo", "ID eliminado", "Texto eliminado (literal)", "Contenido en (ID conservado)",
                   "Texto conservado (literal)", "Razón"])
        for g, rid, cont, raz in eliminados:
            wd.append([g, rid, texto_de[rid], cont, texto_de[cont], raz])
        for col, w in zip("ABCDEFG", (12, 14, 70, 16, 70, 50)):
            wd.column_dimensions[col].width = w
        for row in wd.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        wd.freeze_panes = "A2"
        wd.auto_filter.ref = f"A1:F{wd.max_row}"

    wc = wb.create_sheet("Taxonomía")
    wc.append(["Faceta (TAXONOMIA.md, decisión 10)", "Valores aprobados"])
    for c, v in CATALOGO:
        wc.append([c, v])
    wc.column_dimensions["A"].width = 30
    wc.column_dimensions["B"].width = 140
    for r in wc.iter_rows(min_row=1):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r[0].font = bold

    we = wb.create_sheet("Líneas excluidas")
    we.append(["Documento", "Línea", "Texto de la línea", "Motivo de exclusión"])
    for arch, ln, tx, mot in con.execute("select archivo, linea, texto, motivo from registros_exclusion order by archivo, linea"):
        we.append([arch, ln, tx, mot])
    for col, w in zip("ABCD", (60, 8, 90, 70)):
        we.column_dimensions[col].width = w
    we.freeze_panes = "A2"
    we.auto_filter.ref = f"A1:D{we.max_row}"

    wf = wb.create_sheet("Fuentes")
    wf.append(["Documento", "Registros", "Párrafos extraídos", "Líneas excluidas con motivo"])
    por_doc = {}
    for rid, fs in fuentes.items():
        for f in fs:
            d = por_doc.setdefault(f[0], {"regs": set(), "tot": f[5]})
            d["regs"].add(rid)
    exc = dict(con.execute("select archivo, count(*) from registros_exclusion group by archivo").fetchall())
    for doc in sorted(set(por_doc) | set(exc)):
        d = por_doc.get(doc, {"regs": set(), "tot": 0})
        wf.append([doc, len(d["regs"]), d["tot"], exc.get(doc, 0)])
    wf.append([])
    wf.append(["Fuentes NO incluidas", "Motivo"])
    import importlib.util
    spec = importlib.util.spec_from_file_location("ext", APD / "tools/extraer_registros.py")
    ext = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ext)
    for f, m in ext.EXCLUIDAS:
        wf.append([f, m])
    wf.column_dimensions["A"].width = 80
    wf.column_dimensions["B"].width = 90
    wb.save(a.xlsx)
    print(f"{len(regs)} registros · {len(eliminados)} duplicados eliminados · {len(CATALOGO)} columnas de catálogo · {a.xlsx}")


if __name__ == "__main__":
    main()
