"""Exporta la tabla de catalogación: columna A = registro completo; columnas siguientes = etiquetas del
catálogo (vacías); al final, la ubicación exacta en la fuente para corregir a mano.

Lee las tablas registros / registros_fuente / registros_exclusion que escribe extraer_registros.py.
Uso: python3 tools/exportar_catalogo.py --xlsx salida.xlsx [--db data/rules.sqlite]
"""

from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

APD = Path(__file__).resolve().parents[1]

# Columnas de catálogo, en el orden de decisión dictado por el usuario; valores posibles = propuesta a aprobar.
CATALOGO = [
    # --- 1ª capa: el sistema de clasificación que ya define ai-production-director/SKILL.md ---
    ("Etapa (director §3)", "0 Brand Lock · 1 Creative Strategy · 2 Screenwriting · 3 Cinematic Direction · 4 Shot Planning · "
                            "5 Anchor Images · 6 Video Prompts · 7 Package & Delivery · Post-producción (production-package §3) · "
                            "Gate final (aurora-prompt-linter) · Director (orquestación, todas las etapas)"),
    ("Sub-skill con autoridad (director §1)", "brand-lock-extractor · creative-strategy.md · screenwriter · video/dramaturgia · "
                            "storyboard-architect · ai-video-storyboard · image · visual-prompt-forge · visual-asset-critic · "
                            "video/archivos de modelo · storyboard-html-preview · visual-media (solo §7) · aurora-prompt-linter · "
                            "produccion-visual-sw30 · regímenes del repo · decisiones del usuario"),
    ("Autoridad sobre (director §1)", "Parámetros de marca · Estrategia creativa · Estructura narrativa y diálogo · Dramaturgia y "
                            "lenguaje de cámara · Fuente de verdad estructural · Shot list EXPRESS · Sintaxis final de prompts de imagen · "
                            "Estructura shots→prompt y loop de revisión · Aceptar/rechazar renders · Sintaxis final de prompts de video · "
                            "Formato de entrega visual · Proyectos de animación / material didáctico ES · Veto final sobre prompts"),
    ("Track (director §2)", "EXPRESS · STANDARD · FILM · todos"),
    ("Gate de la etapa (director §3)", "define o verifica el gate de su etapa · no"),
    ("Precedencia en conflicto (director §6)", "6.1 vocabulario prohibido gana · 6.2 sintaxis final del modelo destino · "
                            "6.3 shots.json fuente de verdad · 6.4 forge = flujo de datos, smixs = texto final · 6.5 atribución · no aplica"),
    # --- 2ª capa: subprocesos (Mapa del Spot) y dimensiones dictadas por el usuario ---
    ("Subproceso", "E0.1–E0.4 · E1.1–E1.5 · E2.1–E2.4 · E3.1–E3.5 · E4.1–E4.6 · E5.1–E5.9 · E6.1–E6.6 · E7.1–E7.4 · POST.1–POST.2"),
    ("Medio", "imagen · video · ambos · documento de preproducción (texto) · proceso"),
    ("Tipo de creación (imagen)", "génesis sin referencias · maestro derivado con referencia (T2, variantes) · cuadro compuesto con "
                                  "referencias (T3/T4, producto en escena) · edición quirúrgica (T5) · keyframe FF/LF · no aplica"),
    ("Modo del clip (video)", "texto a video · imagen a video FF · FF+LF · referencias/elements/ingredients · motion control / video de "
                              "movimiento · multi-shot · diálogo y lip-sync · edición de video · extensión · ultra long / blockout · no aplica"),
    ("Tipo de acción (D2)", "A pose sostenida · B instante pico · C movimiento continuo · D fenómeno de luz o clima · "
                            "E interacción o diálogo · F multitud · ensamble · sin acción (gráfico/UI/slide) · no aplica"),
    ("Régimen físico", "01 agua superficie · 02 ruptura de superficie · 03 subacuático · 04 objeto balístico · 05 vehículo · "
                       "06 cuerpo en esfuerzo · 07 luz y clima · 08 multitud anónima · 09 grupo ensamble · no aplica"),
    ("Clase de sujeto", "persona · lugar genérico · edificación/landmark · animal · producto · objeto/prop · gráfico/tipografía/UI · no aplica"),
    ("Espacio", "interior · exterior · estudio o fondo vacío · no aplica"),
    ("Tratamiento (D4)", "documental · narrativo cinematográfico · comercial pulido · race/kinetic · UGC/social · editorial/moda · "
                         "animación/ilustración · gráfico/diseño · no aplica"),
    ("Número de sujetos (D3)", "0 placa · 1 · 2 sin contacto · 2 con contacto · ensamble 3–8 con identidad · multitud anónima · no aplica"),
    ("Rostro en cuadro", "rostro ≥20% del cuadro · rostro <20% · fragmentado/silueta/casco · sin rostro · no aplica"),
    ("Referencias", "sí · no · no aplica"),
    ("Tipo de referencia", "persona (P) · vestuario (O) · lugar (L) · producto/prop (PR) · estilo (S) · frame FF/LF · video de movimiento · "
                           "audio/voz · layout/boceto/grid · no aplica"),
    ("Rol del ancla (D1)", "character ref · environment ref · keyframe FF · keyframe LF · hero/macro · motif insert · no aplica"),
    ("Modelo (D8)", "Nano Banana 2 · Nano Banana Pro · GPT Image 2 · Flux · Midjourney · Ideogram · Seedream · Kling · Veo · "
                    "Seedance 2.0 · Seedance 2.5 · Hailuo · agnóstico"),
    ("Ángulo y altura (D9)", "eye-level · high · low · overhead · dutch · worm's eye · POV · no aplica"),
    ("Cámara y lente (D6)", "encuadre · lente · DOF · movimiento/cámara implícita · rig · no aplica"),
    ("Luz (D7)", "fuente motivada · dirección · dureza · ratio · specular · atmósfera · no aplica"),
    ("Paleta y grade (D5)", "brand-lock hex · series_lock grade · paleta nombrada · no aplica"),
    ("Audio", "diálogo · VO · SFX · música · silencio · no aplica"),
    ("Texto en pantalla", "overlay compuesto en post · texto dentro de la imagen · sin texto · no aplica"),
    ("Formato y plataforma", "9:16 social · 16:9 · 1:1 · 4:5 · póster/impreso · slide · UI/app · no aplica"),
    ("Marca", "con brand-lock · logo o producto real · sin marca · no aplica"),
    ("Tipo de registro", "prohibición · obligación · recomendación · sintaxis/plantilla · límite o dato de plataforma · gate/verificación · "
                         "ejemplo · explicación"),
    ("Ámbito", "pipeline · ejemplo de marca · herramienta/validador · conducta del agente · aprendizaje de producción · "
               "decisión del usuario · candidata a prueba"),
]

UBICACION = ["Documento", "Página / sección", "Párrafo nº (de N)", "Líneas", "ID registro"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default=str(APD / "data/rules.sqlite"))
    ap.add_argument("--xlsx", required=True)
    a = ap.parse_args()
    con = sqlite3.connect(a.db)
    regs = con.execute("select id, orden, texto from registros order by orden").fetchall()
    fuentes: dict[str, list] = {}
    for rid, arch, li, lf, sec, par, tot in con.execute(
            "select registro_id, archivo, linea_ini, linea_fin, seccion, parrafo, parrafos_en_documento from registros_fuente"):
        fuentes.setdefault(rid, []).append((arch, li, lf, sec, par, tot))

    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    cab = ["Registro (regla o punto de sintaxis completo)"] + [c for c, _ in CATALOGO] + UBICACION
    ws.append(cab)
    for rid, orden, texto in regs:
        fs = fuentes.get(rid, [])
        doc = " | ".join(f[0] for f in fs)
        sec = " | ".join(f[3] or "" for f in fs)
        par = " | ".join(f"{f[4]} de {f[5]}" for f in fs)
        lin = " | ".join((f"{f[1]}–{f[2]}" if f[1] and f[2] != f[1] else str(f[1])) if f[1] else "objeto JSON/Python" for f in fs)
        ws.append([texto] + [""] * len(CATALOGO) + [doc, sec, par, lin, rid])
    bold = Font(bold=True)
    fill_cat = PatternFill("solid", fgColor="FFF2CC")
    fill_ubi = PatternFill("solid", fgColor="DDEBF7")
    for col in range(1, len(cab) + 1):
        c = ws.cell(row=1, column=col)
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if 2 <= col <= 1 + len(CATALOGO):
            c.fill = fill_cat
        elif col > 1 + len(CATALOGO):
            c.fill = fill_ubi
    ws.column_dimensions["A"].width = 110
    for col in range(2, 2 + len(CATALOGO)):
        ws.column_dimensions[get_column_letter(col)].width = 16
    for i, w in enumerate([48, 48, 14, 12, 14]):
        ws.column_dimensions[get_column_letter(2 + len(CATALOGO) + i)].width = w
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{ws.max_row}"

    wc = wb.create_sheet("Catálogo propuesto")
    wc.append(["Columna (etiqueta de catálogo)", "Valores propuestos — pendientes de tu aprobación"])
    for c, v in CATALOGO:
        wc.append([c, v])
    wc.column_dimensions["A"].width = 30
    wc.column_dimensions["B"].width = 140
    for r in wc.iter_rows(min_row=1):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r[0].font = bold

    we = wb.create_sheet("Líneas excluidas")
    we.append(["Documento", "Línea", "Texto de la línea", "Motivo de exclusión"])
    for arch, ln, tx, mot in con.execute("select archivo, linea, texto, motivo from registros_exclusion order by archivo, linea"):
        we.append([arch, ln, tx, mot])
    for col, w in zip("ABCD", (60, 8, 90, 70)):
        we.column_dimensions[col].width = w
    we.freeze_panes = "A2"
    we.auto_filter.ref = f"A1:D{we.max_row}"

    wf = wb.create_sheet("Fuentes")
    wf.append(["Documento", "Registros", "Párrafos extraídos", "Líneas excluidas con motivo"])
    por_doc = {}
    for rid, fs in fuentes.items():
        for f in fs:
            d = por_doc.setdefault(f[0], {"regs": set(), "tot": f[5]})
            d["regs"].add(rid)
    exc = dict(con.execute("select archivo, count(*) from registros_exclusion group by archivo").fetchall())
    for doc in sorted(set(por_doc) | set(exc)):
        d = por_doc.get(doc, {"regs": set(), "tot": 0})
        wf.append([doc, len(d["regs"]), d["tot"], exc.get(doc, 0)])
    wf.append([])
    wf.append(["Fuentes NO incluidas", "Motivo"])
    import importlib.util
    spec = importlib.util.spec_from_file_location("ext", APD / "tools/extraer_registros.py")
    ext = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ext)
    for f, m in ext.EXCLUIDAS:
        wf.append([f, m])
    wf.column_dimensions["A"].width = 80
    wf.column_dimensions["B"].width = 90
    wb.save(a.xlsx)
    print(f"{len(regs)} registros · {len(CATALOGO)} columnas de catálogo · {a.xlsx}")


if __name__ == "__main__":
    main()
