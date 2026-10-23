"""construir.py: verifica la clasificación de gpt-image.md y, solo si pasa, escribe la matriz en Excel.

    python3 construir.py
"""
import importlib.util
import json
import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).resolve().parent
V2 = AQUI.parent
sys.path.insert(0, str(V2))
import matriz as M  # noqa: E402  columnas de la taxonomía, de lo general a lo particular

spec = importlib.util.spec_from_file_location("clasificacion", AQUI / "clasificacion.py")
C = importlib.util.module_from_spec(spec)
spec.loader.exec_module(C)
E = M.E
ACCIONES = ["agregar al prompt", "seleccionar estructura", "fijar parámetro", "validar el prompt",
            "validar contra el ejemplo", "excluir del prompt", "solo traza"]
DESTINOS = [d for d in E["destinos"] if d.startswith(("salida.slots", "salida.parametros", "salida.validaciones",
                                                      "salida.referencias", "salida.aceptacion", "salida.traza",
                                                      "entrada.modelo", "entrada.referencias", "salida.prompt"))]
FACETAS_APLIC = ["P01"] + [f"F{i:02d}" for i in range(1, 30)] + ["P04", "F32", "F33", "F34", "F35"]


def columnas():
    cols = M.columnas()
    k = max(i for i, (g, _) in enumerate(cols) if g == "Uso operativo") + 1
    extra = [("Acción al activarse", a) for a in ACCIONES] + [("Destino", d) for d in DESTINOS]
    return cols[:k] + extra + cols[k:]


def g(f):
    return M.grupo_de(f)


def todos(f):
    return [v["valor"] for v in E["facetas"][f]["valores"]]


def marcas(a):
    m = {("Tipo de registro", "regla o punto"), ("Tipo de enunciado", a["t61"]), ("Objeto regulado", a["objeto"]),
         ("Uso operativo", a["uso"]), ("Idioma", a["idioma"]), ("Origen de evidencia", "Evidencia directa"),
         ("Evidencia de la afirmación", a["evid"]), ("Adopción canónica", "adopción documental")}
    m |= {("Tipo de enunciado §1.12", t) for t in a["t112"]}
    m |= {("Fuerza normativa", x) for x in a["fuerza"]}
    m |= {("Ámbito funcional", x) for x in a["ambito"]}
    m |= {("Acción al activarse", x) for x in a["accion"]}
    m |= {("Destino", x) for x in a["destino"]}
    if a["medio"]:
        m.add(("Medio", a["medio"]))
    if a["unidad"]:
        m.add(("Unidad", a["unidad"]))
    if a["modelo_citado"]:
        m.add(("Estado del modelo citado", a["modelo_citado"]))
    if a["estado_cap"]:
        m.add(("Estado de evidencia de capacidad", a["estado_cap"]))
    if a["dependencia"]:
        m.add(("Dependencia", a["dependencia"]))
    fac = {f: todos(f) for f in FACETAS_APLIC}
    fac.update({"F01": ["5 Imagen"], "F02": a["f02"], "F29": a["f29"]})
    fac.update(a["r"])
    for f, vals in fac.items():
        m |= {(g(f), v) for v in vals}
    m |= {(g("F30"), v) for v in a["f30"]}
    m.add((g("F31"), a["f31"]))
    if not any(f in M.CASO for f in a["r"]):
        m.add(("Universal", "aplica a todo caso"))
    return m


def verificar(unidades, cols):
    err = []
    idx = set(cols)
    ids = [u["id"] for u in unidades]
    con_aplic = {a["id"] for a in C.APLICACIONES}
    for i in ids:
        n = (i in C.CONCILIACIONES) + (i in con_aplic)
        if n != 1:
            err.append(f"{i}: debe estar exactamente en conciliaciones o en aplicaciones (está en {n})")
    for i in set(C.CONCILIACIONES) | con_aplic:
        if i not in ids:
            err.append(f"{i}: no pertenece a gpt-image.md")
    tipos_con = {v["valor"] for v in E["campos"]["tipo_conciliacion"]}
    for i, (t, j) in C.CONCILIACIONES.items():
        if t not in tipos_con or len(j.split()) < 4:
            err.append(f"{i}: conciliación sin tipo válido o sin justificación")
    grupos_oblig = ["Tipo de enunciado", "Fuerza normativa", "Ámbito funcional", "Objeto regulado", "Uso operativo",
                    "Acción al activarse", "Destino", "Idioma", "Origen de evidencia", "Evidencia de la afirmación",
                    "Adopción canónica"] + [g(f) for f in FACETAS_APLIC + ["F30", "F31"]]
    vistos = set()
    for a in C.APLICACIONES:
        clave = (a["id"], a["sufijo"])
        if clave in vistos:
            err.append(f"{clave}: aplicación repetida")
        vistos.add(clave)
        mk = marcas(a)
        for c in mk:
            if c not in idx:
                err.append(f"{a['id']}{a['sufijo']}: marca que no existe en la taxonomía: {c}")
        presentes = {gr for gr, _ in mk}
        for gr in grupos_oblig:
            if gr not in presentes:
                err.append(f"{a['id']}{a['sufijo']}: grupo sin ninguna X: {gr}")
        if (a["objeto"] == "activo visual o sonoro") != bool(a["medio"]):
            err.append(f"{a['id']}{a['sufijo']}: medio solo y siempre en activo visual o sonoro")
        if not a["unidad"] and "U1" not in (a["consulta"] or ""):
            err.append(f"{a['id']}{a['sufijo']}: sin unidad y sin consulta")
        restringe = a["r"] or a["f02"] != C.PASOS_PROMPT or a["f29"] != C.GPT2
        if restringe and len(a["fund"].split()) < 3:
            err.append(f"{a['id']}{a['sufijo']}: restricción sin fundamento")
        if a["valor"] in ("", None):
            err.append(f"{a['id']}{a['sufijo']}: falta el valor a insertar o «—»")
        for q in (a["consulta"] or "").replace(" ", "").split("·"):
            if q and q not in C.CONSULTAS:
                err.append(f"{a['id']}{a['sufijo']}: consulta {q} sin texto")
    return err


def main():
    lotes = [json.loads((V2 / "lotes" / f"lote_{n}.json").read_text()) for n in ("09", "10")]
    unidades = [u for lote in lotes for u in lote]
    assert {u["archivo"] for u in unidades} == {"skills/image/references/gpt-image.md"} and len(unidades) == 84
    cols = columnas()
    err = verificar(unidades, cols)
    reglas = len({a["id"] for a in C.APLICACIONES})
    informe = [f"Textos de gpt-image.md: {len(unidades)} (esperados 84)",
               f"Conciliaciones (estructura y no normativos): {len(C.CONCILIACIONES)}",
               f"Textos que son regla o punto: {reglas} · aplicaciones (filas): {len(C.APLICACIONES)}",
               f"Columnas: {len(cols)} (todas generadas desde la taxonomía, más acción y destino)",
               f"Consultas abiertas para Eric: {len(C.CONSULTAS)}",
               "VERIFICACIÓN: " + ("CORRECTA — sin omisiones ni valores fuera de la taxonomía" if not err else
                                   f"{len(err)} ERRORES")] + err
    print("\n".join(informe))
    if err:
        sys.exit(1)
    db = sqlite3.connect(M.BASE)
    tx = {i: (l, t) for i, l, t in db.execute(
        "SELECT id, linea_inicial, texto FROM textos WHERE archivo='skills/image/references/gpt-image.md'")}
    wb = Workbook()
    ws = wb.active
    ws.title = "matriz"
    fijas = ["id", "aplicación", "línea", "texto original", "valor a insertar", "fundamento", "consulta"]
    for k, h in enumerate(fijas, 1):
        ws.cell(row=2, column=k, value=h).font = Font(bold=True)
    previo = None
    gris = PatternFill("solid", fgColor="DDDDDD")
    for k, (gr, v) in enumerate(cols, len(fijas) + 1):
        if gr != previo:
            ws.cell(row=1, column=k, value=gr).font = Font(bold=True)
            previo = gr
        c = ws.cell(row=2, column=k, value=v)
        c.alignment = Alignment(text_rotation=90, vertical="bottom")
        c.fill = gris
    pos = {c: k for k, c in enumerate(cols, len(fijas) + 1)}
    filas = []
    for u in unidades:
        if u["id"] in C.CONCILIACIONES:
            t, j = C.CONCILIACIONES[u["id"]]
            filas.append((u["id"], "—", "—", j, "", {("Tipo de registro", t)}))
        else:
            for a in [x for x in C.APLICACIONES if x["id"] == u["id"]]:
                filas.append((u["id"], a["sufijo"] or "A1", a["valor"], a["fund"], a["consulta"] or "", marcas(a)))
    for n, (i, suf, valor, fund, cons, mk) in enumerate(filas, 3):
        linea, texto = tx[i]
        for k, v in enumerate([i, suf, linea, texto, valor, fund, cons], 1):
            ws.cell(row=n, column=k, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        for c in mk:
            ws.cell(row=n, column=pos[c], value="X").alignment = Alignment(horizontal="center")
    for col, w in zip("ABCDEFG", (11, 6, 6, 60, 40, 40, 9)):
        ws.column_dimensions[col].width = w
    for k in range(len(fijas) + 1, len(fijas) + len(cols) + 1):
        ws.column_dimensions[get_column_letter(k)].width = 3.2
    ws.row_dimensions[2].height = 200
    ws.freeze_panes = "H3"
    hc = wb.create_sheet("criterios")
    for linea in (C.__doc__ or "").strip().splitlines():
        hc.append([linea])
    hq = wb.create_sheet("consultas")
    hq.append(["código", "pregunta", "filas"])
    for q, texto in C.CONSULTAS.items():
        hq.append([q, texto, ", ".join(sorted({f"{a['id']}{a['sufijo']}" for a in C.APLICACIONES
                                               if q in (a["consulta"] or "")}))])
    hv = wb.create_sheet("verificación")
    for linea in informe:
        hv.append([linea])
    hcol = wb.create_sheet("columnas")
    hcol.append(["n", "grupo", "valor"])
    for k, (gr, v) in enumerate(cols, 1):
        hcol.append([k, gr, v])
    salida = AQUI / "matriz_gpt-image_desde_cero.xlsx"
    wb.save(salida)
    print(f"escrito {salida.name}: {len(filas)} filas × {len(fijas) + len(cols)} columnas")
    # Comprobación del texto: se reabre el Excel y cada celda se compara con la base original, carácter por
    # carácter y por sha256 (el mismo hash que certifica cada unidad en la base de 13,005 textos).
    import hashlib
    from openpyxl import load_workbook
    sha = {i: h for i, h in db.execute("SELECT id, sha256_unidad FROM textos WHERE "
                                       "archivo='skills/image/references/gpt-image.md'")}
    leido = load_workbook(salida)["matriz"]
    iguales, distintos, vistos = 0, [], set()
    for fila in leido.iter_rows(min_row=3, values_only=True):
        i, texto = fila[0], fila[3]
        vistos.add(i)
        ok = texto == tx[i][1] and hashlib.sha256(texto.encode("utf-8")).hexdigest() == sha[i]
        iguales += ok
        if not ok:
            distintos.append(i)
    faltan = sorted(set(tx) - vistos)
    print(f"texto literal: {iguales} filas idénticas al original (carácter y sha256) · distintas {distintos} · "
          f"textos del archivo que faltan en la matriz {faltan}")
    if distintos or faltan:
        sys.exit(1)


if __name__ == "__main__":
    main()
