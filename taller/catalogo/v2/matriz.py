"""matriz.py: matriz texto × clasificación (X o en blanco), columnas de lo general a lo particular.

Las columnas salen de esquema_taxonomia.json (taxonomía original + extensiones aprobadas): no puede haber una X en
un valor que no exista. Una faceta UNIVERSAL marca X en todos sus valores. La columna "Universal" se calcula: X si
el texto es regla y no restringe ninguna faceta de caso (F03–F28, P04, F32–F34).

    python3 matriz.py <salida_NN.json>... --xlsx matriz.xlsx
"""
import argparse
import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).resolve().parent
E = json.loads((AQUI / "esquema_taxonomia.json").read_text())
CASO = [f"F{i:02d}" for i in range(3, 29)] + ["P04", "F32", "F33", "F34"]
T112 = json.loads((AQUI / "tipos_enunciado_112.json").read_text())
TIPOS_112 = list(T112["criterio"])
assert all(t in T112["fuente"] for t in TIPOS_112), "cada tipo debe estar literal en la línea 192 citada"
BASE = ("/tmp/claude-0/-home-user-AI-Production-Director/5525717f-6970-5668-b1ce-508e1e3863a0/scratchpad/docs/"
        "d4ea0f84-BASE_TEXTOS_ORIGINALES_13005_20261019_055700_1/BASE_TEXTOS_ORIGINALES_13005_20261019_055700.sqlite")


def columnas():
    """(grupo, valor) en orden de lo general a lo particular."""
    cols = [("Tipo de registro", v) for v in ["regla o punto"] + [x["valor"] for x in E["campos"]["tipo_conciliacion"]]]
    cols += [("Tipo de enunciado §1.12", t) for t in TIPOS_112]
    for campo, nombre in (("tipo_enunciado", "Tipo de enunciado"), ("fuerza_normativa", "Fuerza normativa"),
                          ("ambito_funcional", "Ámbito funcional"), ("objeto_regulado", "Objeto regulado"),
                          ("medio", "Medio"), ("unidad", "Unidad"), ("uso_operativo", "Uso operativo")):
        cols += [(nombre, x["valor"]) for x in E["campos"][campo]]
    cols += [(f"F29 {E['facetas']['F29']['nombre']}", x["valor"]) for x in E["facetas"]["F29"]["valores"]]
    cols.append(("Universal", "aplica a todo caso"))
    orden = ["F01", "F02", "P01", "F03", "F04", "P04", "F05", "F06", "F07", "F08", "F09", "F10", "F11", "F12", "F13",
             "F14", "F15", "F16", "F17", "F18", "F19", "F20", "F21", "F22", "F23", "F24", "F25", "F26", "F27", "F28",
             "F32", "F33", "F34", "F30", "F31"]
    for f in orden:
        cols += [(f"{f} {E['facetas'][f]['nombre']}", x["valor"]) for x in E["facetas"][f]["valores"]]
    for campo, nombre in (("idioma", "Idioma"), ("origen_evidencia", "Origen de evidencia"),
                          ("evidencia_afirmacion_fuente", "Evidencia de la afirmación"),
                          ("adopcion_canonica", "Adopción canónica"), ("estado_modelo_citado", "Estado del modelo citado"),
                          ("estado_evidencia", "Estado de evidencia de capacidad"), ("dependencia", "Dependencia")):
        cols += [(nombre, x["valor"]) for x in E["campos"][campo]]
    return cols


def grupo_de(f):
    return f"{f} {E['facetas'][f]['nombre']}"


def marcas(r):
    """Conjunto de (grupo, valor) con X para un registro."""
    if r["tipo_registro"] == "conciliacion":
        return {("Tipo de registro", r["tipo_conciliacion"])}
    m = {("Tipo de registro", "regla o punto")}
    m |= {("Tipo de enunciado §1.12", t) for t in T112["asignaciones"].get(r["id"], [])}
    be, bd = r["bloque_E"], r["bloque_D"]
    for campo, nombre in (("tipo_enunciado", "Tipo de enunciado"), ("fuerza_normativa", "Fuerza normativa"),
                          ("objeto_regulado", "Objeto regulado"), ("medio", "Medio"), ("unidad", "Unidad"),
                          ("idioma", "Idioma"), ("origen_evidencia", "Origen de evidencia"),
                          ("evidencia_afirmacion_fuente", "Evidencia de la afirmación"),
                          ("adopcion_canonica", "Adopción canónica"), ("estado_evidencia", "Estado de evidencia de capacidad"),
                          ("dependencia", "Dependencia")):
        if be[campo] is not None:
            m.add((nombre, be[campo]))
    if be["estado_modelo_citado"]:
        m.add(("Estado del modelo citado", be["estado_modelo_citado"]["estado"]))
    m |= {("Ámbito funcional", a) for a in bd["ambito_funcional"]}
    for f in ("F30", "F31"):
        m |= {(grupo_de(f), v) for v in bd[f]["valores"]}
    restringe_caso = False
    for ap in r["aplicaciones"]:
        uso = ap["uso_operativo"]
        if uso == "INCORPORADA A TEMPLATE":  # §5.3: sin template verificado todavía es construcción de template
            uso = "APLICACIÓN EN CONSTRUCCIÓN DE TEMPLATE"
        m.add(("Uso operativo", uso))
        for f, v in ap["facetas"].items():
            vals = [x["valor"] for x in E["facetas"][f]["valores"]] if v["estado"] == "UNIVERSAL" else v["valores"]
            m |= {(grupo_de(f), x) for x in vals}
            restringe_caso |= f in CASO and v["estado"] == "VALOR"
    if not restringe_caso:
        m.add(("Universal", "aplica a todo caso"))
    return m


def textos(ids):
    db = sqlite3.connect(BASE)
    q = f"SELECT id, archivo, linea_inicial, texto FROM textos WHERE id IN ({','.join('?' * len(ids))})"
    return {i: (a, l, t) for i, a, l, t in db.execute(q, ids)}


def construir(salidas, ruta_xlsx):
    cols = columnas()
    idx = {c: k for k, c in enumerate(cols)}
    filas = []
    for s in salidas:
        for r in json.loads(Path(s).read_text()):
            mk = marcas(r)
            fuera = [c for c in mk if c not in idx]
            assert not fuera, f"{r['id']}: marcas fuera de la taxonomía {fuera}"
            for u in r["unidades"]:
                filas.append((u, r["id"], mk))
    tx = textos([f[0] for f in filas])
    wb = Workbook()
    ws = wb.active
    ws.title = "matriz"
    fijas = ["id", "regla (id)", "archivo:línea", "texto"]
    for k, h in enumerate(fijas, 1):
        ws.cell(row=2, column=k, value=h).font = Font(bold=True)
    gris = PatternFill("solid", fgColor="DDDDDD")
    previo = None
    for k, (g, v) in enumerate(cols, len(fijas) + 1):
        if g != previo:
            ws.cell(row=1, column=k, value=g).font = Font(bold=True)
            previo = g
        c = ws.cell(row=2, column=k, value=v)
        c.alignment = Alignment(text_rotation=90, vertical="bottom")
        c.fill = gris
    for n, (u, rid, mk) in enumerate(filas, 3):
        a, l, t = tx[u]
        ws.cell(row=n, column=1, value=u)
        ws.cell(row=n, column=2, value=rid)
        ws.cell(row=n, column=3, value=f"{a.split('/')[-1]}:{l}")
        ws.cell(row=n, column=4, value=t[:1500]).alignment = Alignment(wrap_text=True, vertical="top")
        for c in mk:
            ws.cell(row=n, column=len(fijas) + 1 + idx[c], value="X").alignment = Alignment(horizontal="center")
    ws.column_dimensions["D"].width = 70
    for k in range(len(fijas) + 1, len(fijas) + len(cols) + 1):
        ws.column_dimensions[get_column_letter(k)].width = 3.2
    ws.row_dimensions[2].height = 190
    ws.freeze_panes = "E3"
    hc = wb.create_sheet("columnas")
    hc.append(["n", "grupo", "valor", "origen"])
    origen = {(f"{f} {d['nombre']}", v["valor"]): v["origen"] for f, d in E["facetas"].items() for v in d["valores"]}
    for k, (g, v) in enumerate(cols, 1):
        hc.append([k, g, v, origen.get((g, v), "taxonomia_original" if g not in ("Tipo de registro", "Universal")
                                       else ("taxonomia_original" if v != "regla o punto" else "estructura de la matriz"))])
    wb.save(ruta_xlsx)
    print(f"{ruta_xlsx}: {len(filas)} textos × {len(cols)} columnas · X marcadas {sum(len(f[2]) for f in filas)}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("salidas", nargs="+")
    ap.add_argument("--xlsx", required=True)
    a = ap.parse_args()
    construir(a.salidas, a.xlsx)
