"""matriz_eric.py: matriz en el formato pedido por Eric (2026-10-24), para cualquier archivo clasificado.

- Sin textos en la matriz (instrucción de Eric): columna A = número del texto (id persistente de la base de 13,005
  textos originales), columna B = línea del archivo. El texto se localiza en la base por ese número.
- Una fila por texto. Si el texto tiene casos alternativos (§7.2), sus X van en la misma fila y la columna "casos"
  dice qué valores se combinan en cada caso.
- Columnas de lo general a lo particular: tipo de texto → imagen o video → modelo → universal → interno o externo →
  tipo de pieza → sujeto → número de personas → resto de facetas → proceso → atributos → acción, destino, valor
  a insertar y fundamento al final.
- Antes de escribir corre la verificación de la clasificación; después reabre el Excel y comprueba el formato y que
  los números sean exactamente los textos del archivo en la base, sin faltantes ni repetidos. Si algo falla, sale
  con error.

    python3 matriz_eric.py <carpeta_con_clasificacion.py> <archivo_fuente>
"""
import importlib.util
import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

V2 = Path(__file__).resolve().parent
sys.path.insert(0, str(V2))
sys.path.insert(0, str(V2 / "piloto_gpt-image"))
import matriz as M  # noqa: E402
import construir as K  # noqa: E402  marcas(), verificar(), ACCIONES, DESTINOS, FACETAS_APLIC

ORDEN = (
    [("Tipo de registro", None), ("Tipo de enunciado §1.12", None), ("Tipo de enunciado", None), ("Fuerza normativa", None),
     ("Ámbito funcional", None), ("Medio", None), ("F29", "f"), ("Universal", None), ("Objeto regulado", None),
     ("Uso operativo", None),
     ("F04", "f"), ("P04", "f"), ("F05", "f"), ("F06", "f"),
     ("F15", "f"), ("F19", "f"), ("F20", "f"), ("F21", "f"), ("F22", "f"), ("F23", "f"),
     ("F10", "f"), ("F13", "f"), ("F14", "f"), ("F16", "f"), ("F17", "f"), ("F18", "f"), ("F24", "f"), ("F25", "f"),
     ("F32", "f"), ("F33", "f"), ("F34", "f"), ("F35", "f"),
     ("F11", "f"), ("F12", "f"), ("F03", "f"), ("F07", "f"), ("F08", "f"), ("F09", "f"), ("F26", "f"), ("F27", "f"),
     ("F28", "f"),
     ("F01", "f"), ("F02", "f"), ("P01", "f"),
     ("F30", "f"), ("F31", "f"), ("Unidad", None), ("Idioma", None), ("Origen de evidencia", None),
     ("Evidencia de la afirmación", None), ("Adopción canónica", None), ("Estado del modelo citado", None),
     ("Estado de evidencia de capacidad", None), ("Dependencia", None),
     ("Acción al activarse", None), ("Destino", None)])


def columnas():
    base = K.columnas()
    por_grupo = {}
    for gr, v in base:
        por_grupo.setdefault(gr, []).append(v)
    cols = []
    for clave, tipo in ORDEN:
        gr = M.grupo_de(clave) if tipo == "f" else clave
        cols += [(gr, v) for v in por_grupo.pop(gr)]
    assert not por_grupo, f"grupos sin ubicar en el orden: {list(por_grupo)}"
    return cols


def describir_caso(a):
    return " + ".join(f"{f}: {', '.join(v)}" for f, v in a["r"].items()) or "sin restricción de caso"


def main(carpeta, archivo):
    spec = importlib.util.spec_from_file_location("clasificacion", Path(carpeta) / "clasificacion.py")
    C = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(C)
    K.C = C  # el verificador usa esta clasificación
    db = sqlite3.connect(M.BASE)
    filas_base = db.execute("SELECT id, linea_inicial, texto, sha256_unidad FROM textos WHERE archivo=? "
                            "ORDER BY CAST(linea_inicial AS INTEGER), id", (archivo,)).fetchall()
    unidades = [{"id": i} for i, *_ in filas_base]
    cols = columnas()
    err = K.verificar(unidades, cols)
    print(f"{archivo}: {len(unidades)} textos · verificación: {'CORRECTA' if not err else f'{len(err)} ERRORES'}")
    for e in err:
        print("  ", e)
    if err:
        sys.exit(1)
    wb = Workbook()
    ws = wb.active
    ws.title = "matriz"
    fijas_ini = ["número", "línea"]
    fijas_fin = ["casos", "valor a insertar", "fundamento", "consulta"]
    for k, h in enumerate(fijas_ini, 1):
        ws.cell(row=2, column=k, value=h).font = Font(bold=True)
    gris = PatternFill("solid", fgColor="DDDDDD")
    previo = None
    pos = {}
    for k, (gr, v) in enumerate(cols, len(fijas_ini) + 1):
        if gr != previo:
            ws.cell(row=1, column=k, value=gr).font = Font(bold=True)
            previo = gr
        c = ws.cell(row=2, column=k, value=v)
        c.alignment = Alignment(text_rotation=90, vertical="bottom")
        c.fill = gris
        pos[(gr, v)] = k
    fin0 = len(fijas_ini) + len(cols) + 1
    for k, h in enumerate(fijas_fin, fin0):
        ws.cell(row=2, column=k, value=h).font = Font(bold=True)
    for n, (i, linea, texto, _) in enumerate(filas_base, 3):
        ws.cell(row=n, column=1, value=i)
        ws.cell(row=n, column=2, value=int(linea))
        if i in C.CONCILIACIONES:
            t, j = C.CONCILIACIONES[i]
            mk, extra = {("Tipo de registro", t)}, ["", "", j, ""]
        else:
            apls = [a for a in C.APLICACIONES if a["id"] == i]
            mk = set().union(*(K.marcas(a) for a in apls))
            casos = " | ".join(f"{a['sufijo'] or 'A1'}: {describir_caso(a)}" for a in apls) if len(apls) > 1 else ""
            extra = [casos, " | ".join(a["valor"] for a in apls), " | ".join(a["fund"] for a in apls),
                     " · ".join(sorted({a["consulta"] for a in apls if a["consulta"]}))]
        for c in mk:
            ws.cell(row=n, column=pos[c], value="X").alignment = Alignment(horizontal="center")
        for k, v in enumerate(extra, fin0):
            ws.cell(row=n, column=k, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 7
    for k in range(len(fijas_ini) + 1, fin0):
        ws.column_dimensions[get_column_letter(k)].width = 3.2
    for k, w in zip(range(fin0, fin0 + 4), (40, 50, 50, 10)):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.row_dimensions[2].height = 200
    ws.freeze_panes = "C3"
    hc = wb.create_sheet("criterios")
    for linea in ((C.__doc__ or "") + "\n" + (__doc__ or "")).strip().splitlines():
        hc.append([linea])
    hq = wb.create_sheet("consultas")
    hq.append(["código", "pregunta"])
    for q, t in C.CONSULTAS.items():
        hq.append([q, t])
    hcol = wb.create_sheet("columnas")
    hcol.append(["n", "grupo", "valor"])
    for k, (gr, v) in enumerate(cols, 1):
        hcol.append([k, gr, v])
    salida = Path(carpeta) / f"matriz_{Path(archivo).stem}.xlsx"
    wb.save(salida)
    leido = load_workbook(salida)["matriz"]
    # Formato pedido por Eric: el programa lo comprueba igual que el contenido; si no se cumple, sale con error.
    formato = []
    if leido.cell(row=2, column=1).value != "número" or leido.cell(row=2, column=2).value != "línea":
        formato.append("la columna A no es el número o la B no es la línea")
    if leido.max_column != fin0 + 3:
        formato.append("hay columnas de más o de menos")
    grupos_leidos = [leido.cell(row=1, column=k).value for k in range(len(fijas_ini) + 1, fin0)
                     if leido.cell(row=1, column=k).value]
    grupos_esperados = list(dict.fromkeys(gr for gr, _ in cols))
    if grupos_leidos != grupos_esperados:
        formato.append("el orden de los grupos de columnas no es el de lo general a lo particular definido en ORDEN")
    nums = [fila[0] for fila in leido.iter_rows(min_row=3, values_only=True)]
    if len(nums) != len(set(nums)):
        formato.append("hay textos repetidos en más de una fila")
    print("formato: " + ("CORRECTO (número en A, línea en B, sin textos, una fila por texto, columnas de general a particular)"
                         if not formato else " · ".join(formato)))
    if formato:
        sys.exit(1)
    linea_de = {i: int(l) for i, l, _, _ in filas_base}
    malas = [n for n, fila in zip(nums, leido.iter_rows(min_row=3, values_only=True)) if linea_de.get(n) != fila[1]]
    faltan, sobran = sorted(set(linea_de) - set(nums)), sorted(set(nums) - set(linea_de))
    print(f"escrito {salida.name}: {len(nums)} filas (una por texto) × {len(cols) + 6} columnas · números del archivo "
          f"en la base {len(linea_de)} · faltan {faltan} · sobran {sobran} · línea distinta {malas}")
    if faltan or sobran or malas:
        sys.exit(1)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
