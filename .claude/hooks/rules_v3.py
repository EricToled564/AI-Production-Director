#!/usr/bin/env python3
"""Base de reglas v3 del pipeline visual — determinista y reproducible.

Orden de construcción (cada paso es idempotente y se puede repetir solo):

    python3 rules_v3.py build  --db rules.sqlite      # 1. esquema + registro de los 13 skills
                                                       #    + regímenes + import-app + FTS
    python3 rules_v3.py derive --db rules.sqlite      # 2. tareas, casos por ruta y facetas
    <venv>/bin/python rules_v3.py embed --db rules.sqlite   # 3. vectores (fastembed + numpy)
    python3 rules_v3.py stats  --db rules.sqlite      # 4. conteos
    python3 rules_v3.py check  --db rules.sqlite      # 5. exit 1 si algo quedó sin caso NI tarea

Subcomandos sueltos, por si hay que repetir un paso: `import-app` (clasificación v2
de la app, .claude/rules/clasificacion_v2_app.json) y `fts` (reconstruye reglas_fts).
`build` ya los ejecuta; `embed` es el único que necesita fastembed y numpy (el venv).

Qué garantiza:

  * Nada desaparece en silencio. Lo que el extractor excluye va a `descartes`; lo
    que una importación no puede colocar va a `importacion_huerfana`; lo que no
    tiene caso, tarea o faceta aparece en `v_sin_faceta` con lo que le falta.
  * Los ids son los de rule_registry.py: sha1(ruta + texto sin etiqueta), estables.
  * Toda derivación es por ruta o por regex declarada en este archivo; se puede
    discutir línea por línea y se reproduce con el mismo comando.
"""

from __future__ import annotations

import argparse
import fnmatch
import hashlib
import importlib.util
import json
import os
import re
import sqlite3
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

HOOKS = Path(__file__).resolve().parent
REPO = HOOKS.parent.parent
SCHEMA = HOOKS / "schema_v3.sql"
CLASIFICACION_V2 = REPO / ".claude" / "rules" / "clasificacion_v2_app.json"
BUILD_APP = REPO / "app" / "build_app.py"
MODELO_EMB = "intfloat/multilingual-e5-large"
FASTEMBED_CACHE = Path(os.environ.get("FASTEMBED_CACHE_PATH", "/tmp/fastembed_cache"))

sys.path.insert(0, str(HOOKS))
import rule_registry as rr  # noqa: E402

# --------------------------------------------------------------------------- taxonomías

CASOS: dict[str, str] = {
    "T1": "maestro de rostro (still, persona, identidad)",
    "T2": "maestro de cuerpo (still, persona completa)",
    "T3": "cuadro con 1 persona (still de escena)",
    "T4": "cuadro con 2 personas (still, contacto/interaccion)",
    "T5": "edicion quirurgica sobre imagen existente",
    "PROD": "still de producto, comida u objeto comercial (sin persona)",
    "GRAF": "pieza grafica: poster, slide, UI, social, tipografia dominante",
    "MULTI": "multi-panel o grid dentro de una sola imagen",
    "REF": "analisis de imagen de referencia, image-to-prompt",
    "CLIP": "prompt de video / movimiento",
    "SHOT": "planeacion de shots, storyboard, shots.json",
    "GUION": "escritura de guion o tratamiento",
    "MARCA": "extraccion o uso de brand-lock",
    "QA": "critica de render, revision de asset generado",
    "ENTREGA": "empaquetado, preview, entrega a cliente",
    "NINGUNO": "no aplica a ningun caso de produccion (meta, ejemplos, docs)",
    # Nuevos en v3: clases de activo genesis de E5.1 (personas, lugares, edificaciones,
    # animales, productos, objetos) y los dos regimenes de grupo.
    "LUGAR": "placa de lugar: entorno o locacion canonizada sin persona (environment ref)",
    "EDIF": "edificacion o landmark canonizado: arquitectura reconocible como activo",
    "ANIMAL": "animal como sujeto: maestro de animal o animal en escena",
    "OBJ": "objeto o prop canonizado que no es producto comercial: vehiculo, herramienta, pelota",
    "MULTITUD": "multitud anonima como entorno: estadio, publico, calle llena, sin rostros legibles",
    "ENSAMBLE": "grupo con identidad: equipo, banda, familia; cada miembro es una identidad",
}

# Subprocesos por etapa. Fuente: flujo-spot.html, seccion "Nivel 2 — Subprocesos por
# etapa" (sesion 2026-09-25), parseada una vez y fijada aqui para que la base se
# reconstruya desde el repo. `parse_flujo_html()` regenera esta tabla desde el HTML.
TAREAS: list[dict] = [
    {"codigo": "E0.1", "etapa": "E0", "nombre": "Reunir fuentes", "autoridad": "brand-lock-extractor", "entra": "URL, brand book PDF, screenshots o descripción", "sale": "Material leído, una sola pregunta si falta todo", "gate": "Paso 1 del workflow"},
    {"codigo": "E0.2", "etapa": "E0", "nombre": "Extraer las 9 secciones", "autoridad": "brand-lock-extractor", "entra": "Fuentes + extraction-rubric.md", "sale": "Identity, palette, type, mood, never list, ratios, grade, motion, voice", "gate": "Regla 1: nunca inventar un hex"},
    {"codigo": "E0.3", "etapa": "E0", "nombre": "Asignar confianza", "autoridad": "brand-lock-extractor", "entra": "Cada valor con su origen", "sale": "extracted / inferred / needs confirmation", "gate": "Regla 3: todo con fuente"},
    {"codigo": "E0.4", "etapa": "E0", "nombre": "Escribir brand-lock.md + Extraction notes", "autoridad": "brand-lock-extractor", "entra": "Template brand-lock.md.tpl", "sale": "brand-lock.md", "gate": "validate_brand_lock.py"},
    {"codigo": "E1.1", "etapa": "E1", "nombre": "Lectura del brief", "autoridad": "creative-strategy.md", "entra": "Brief + brand-lock", "sale": "Problema, audiencia, single-minded proposition en una frase, tono, restricciones", "gate": "Sin SMP en una frase, la etapa no continúa"},
    {"codigo": "E1.2", "etapa": "E1", "nombre": "Territorios creativos", "autoridad": "creative-strategy.md", "entra": "SMP", "sale": "3 a 4 territorios distintos", "gate": "Prueba de distancia: mismo primer plano = mismo territorio"},
    {"codigo": "E1.3", "etapa": "E1", "nombre": "Concept cards", "autoridad": "creative-strategy.md", "entra": "Territorios fuertes", "sale": "2 a 3 cards en STANDARD, 3 a 5 en FILM, con complejidad AI", "gate": "Formato obligatorio de 8 campos"},
    {"codigo": "E1.4", "etapa": "E1", "nombre": "Matriz de selección", "autoridad": "creative-strategy.md", "entra": "Cards", "sale": "Matriz 1 a 5 + recomendación argumentada", "gate": "Usuario elige UN concepto"},
    {"codigo": "E1.5", "etapa": "E1", "nombre": "Dirección narrativa", "autoridad": "creative-strategy.md", "entra": "Concepto ganador", "sale": "Estructura, POV, 5 anclas preliminares, reglas del mundo, densidad", "gate": "Si screenwriter pregunta algo estructural, se corrige aquí"},
    {"codigo": "E2.1", "etapa": "E2", "nombre": "Analizar concepto", "autoridad": "screenwriter", "entra": "Card + dirección narrativa", "sale": "Beats, personajes, locaciones, arco", "gate": "Core workflow 1"},
    {"codigo": "E2.2", "etapa": "E2", "nombre": "Scene breakdown", "autoridad": "screenwriter", "entra": "Beats", "sale": "Escenas con propósito y progresión emocional", "gate": "Cada escena cambia emoción, avanza acción o sube presión"},
    {"codigo": "E2.3", "etapa": "E2", "nombre": "Escribir escenas", "autoridad": "screenwriter", "entra": "Breakdown", "sale": "Sluglines, acción visual, diálogo esencial, intro de personajes", "gate": "Sin direcciones de cámara, sin emociones nombradas"},
    {"codigo": "E2.4", "etapa": "E2", "nombre": "Salida XML-tagged", "autoridad": "screenwriter", "entra": "Escenas", "sale": "script.md con key_visuals por escena", "gate": "Gate E2: fórmula de escena de dramaturgy.md"},
    {"codigo": "E3.1", "etapa": "E3", "nombre": "Fórmula de escena", "autoridad": "video/dramaturgy.md §1", "entra": "Cada escena del script", "sale": "Deseo, obstáculo, geometría, mirada, ritmo nombrados en una frase", "gate": "Si falta uno, la escena no está lista"},
    {"codigo": "E3.2", "etapa": "E3", "nombre": "Blocking y staging", "autoridad": "video/dramaturgy.md §5 y §6", "entra": "Escena", "sale": "Quién quiere qué de quién; dinámica de poder en el encuadre", "gate": "Nombrar la dinámica antes de escribir"},
    {"codigo": "E3.3", "etapa": "E3", "nombre": "Cámara motivada y lente", "autoridad": "video/dramaturgy.md §7, universal-rules U5 y U6", "entra": "Escena", "sale": "Un movimiento primario por shot con su razón: qué cambió", "gate": "Regla Fincher: sin razón, cámara estática"},
    {"codigo": "E3.4", "etapa": "E3", "nombre": "Luz, color y ritmo de edición", "autoridad": "video/dramaturgy.md §9, §10 y §12", "entra": "Escena", "sale": "Presión ambiental, evolución de color, escalera de ritmo", "gate": "Una presión ambiental por escena"},
    {"codigo": "E3.5", "etapa": "E3", "nombre": "Los 5 anclas definitivos", "autoridad": "video/dramaturgy.md §13", "entra": "Anclas preliminares de E1.5", "sale": "Emoción, motivo, objeto, quiebre, imagen final", "gate": "Gate E3: cero vocabulario prohibido. El objeto ancla es un candidato a génesis"},
    {"codigo": "E4.1", "etapa": "E4", "nombre": "Leer brand-lock y elegir beat framework", "autoridad": "storyboard-architect pasos 1 y 2", "entra": "brand-lock + brief", "sale": "Framework: Pain-reframe-promise, Hero Trilogy, etc.", "gate": "beat-frameworks.md"},
    {"codigo": "E4.2", "etapa": "E4", "nombre": "Timing", "autoridad": "storyboard-architect paso 3", "entra": "Duración total", "sale": "Hook 0–2s, setup, prueba, promesa/CTA", "gate": "timing-rules.md"},
    {"codigo": "E4.3", "etapa": "E4", "nombre": "Shot list", "autoridad": "storyboard-architect paso 4", "entra": "Script + dirección", "sale": "shots[]: id, beat, start/end, framing, angle, motion, subject, refs, vo, rationale", "gate": "shots.schema.json, additionalProperties false"},
    {"codigo": "E4.4", "etapa": "E4", "nombre": "Capa de texto separada", "autoridad": "storyboard-architect paso 5", "entra": "Texto en pantalla", "sale": "text-overlays.json", "gate": "Nunca texto dentro de la descripción visual"},
    {"codigo": "E4.5", "etapa": "E4", "nombre": "Series lock", "autoridad": "storyboard-architect paso 6", "entra": "Anclas de E3", "sale": "character, environment, lighting, color_grade", "gate": "Define qué génesis hacen falta: personaje y entorno como mínimo"},
    {"codigo": "E4.6", "etapa": "E4", "nombre": "Rationale, snapshot y run.json", "autoridad": "storyboard-architect pasos 7 a 9", "entra": "Los archivos finales", "sale": "storyboard.md, brand-lock.snapshot.md, run.json con SHA-256", "gate": "validate_shots.py, validate_brand_lock.py, validate_provenance.py"},
    {"codigo": "E5.1", "etapa": "E5", "nombre": "Plan de anchors", "autoridad": "ai-production-director Etapa 5", "entra": "shots.json + brand-lock", "sale": "Inventario de activos génesis por clase: personas, lugares, edificaciones, animales, productos, objetos; qué shot usa cada uno", "gate": "anchor-plan.md"},
    {"codigo": "E5.2", "etapa": "E5", "nombre": "Génesis sin referencias, una por clase de sujeto", "autoridad": "image + produccion-visual-sw30", "entra": "Solo texto: plan de anchors + brand-lock + series_lock", "sale": "Activo canonizado: maestro de persona, placa de lugar, edificación o landmark, animal, producto, objeto", "gate": "Subflujo atómico, Nivel 4"},
    {"codigo": "E5.3", "etapa": "E5", "nombre": "Maestros derivados con referencia", "autoridad": "image/characters.md + SW30 R5, R6", "entra": "Génesis canonizada", "sale": "Maestro de cuerpo T2, variantes de pose y vestuario, animal en otras poses", "gate": "Herencia mínima: same face and clothes + deltas"},
    {"codigo": "E5.4", "etapa": "E5", "nombre": "Cuadros con referencias", "autoridad": "image + SW30 R2, R7, R9", "entra": "Sujeto canonizado + placa canonizada", "sale": "Cuadro 1 sujeto T3, cuadro 2 sujetos T4, producto en escena", "gate": "Mundos por dos pasos; posición por ancla legible de la imagen"},
    {"codigo": "E5.5", "etapa": "E5", "nombre": "Edición quirúrgica", "autoridad": "image/editing.md + SW30 R3, R12", "entra": "Cuadro aprobado ≥80%", "sale": "Cuadro corregido T5", "gate": "Edit, don't re-roll; hasta 3 ediciones apiladas"},
    {"codigo": "E5.6", "etapa": "E5", "nombre": "Borradores por adaptador", "autoridad": "visual-prompt-forge", "entra": "shots.json validado + adapters/*.md", "sale": "prompts/round-N/<modelo>.txt", "gate": "validate_prompts.py: series_lock verbatim, sin texto, sin logos"},
    {"codigo": "E5.7", "etapa": "E5", "nombre": "Crítica por capas", "autoridad": "visual-asset-critic", "entra": "Render + shot + prompt + brand-lock", "sale": "critique.json: ACCEPT / REVISE / REJECT con severidades", "gate": "validate_critique.py, validate_provenance.py"},
    {"codigo": "E5.8", "etapa": "E5", "nombre": "Revisión selectiva", "autoridad": "visual-prompt-forge modo revisión", "entra": "critique.json de la ronda más alta", "sale": "revised-<modelo>.txt solo con shots REVISE", "gate": "Se detiene en REJECT; máximo 2 rondas"},
    {"codigo": "E5.9", "etapa": "E5", "nombre": "Canonizar", "autoridad": "produccion-visual-sw30 R11", "entra": "Render ACCEPT", "sale": "Asset con hash, referencia para lo siguiente", "gate": "Lo caro nunca se regenera"},
    {"codigo": "E6.1", "etapa": "E6", "nombre": "Motion draft desde shots.json + anchor", "autoridad": "visual-prompt-forge", "entra": "Shot + keyframes start/tail aprobados", "sale": "Borrador por adaptador kling / veo / seedance / hailuo", "gate": "_capabilities.json manda sobre el .md"},
    {"codigo": "E6.2", "etapa": "E6", "nombre": "Reescritura a sintaxis del modelo", "autoridad": "video/kling.md, veo.md o seedance.md", "entra": "Borrador", "sale": "Motion brief: cámara, blocking, performance, luz, física, audio", "gate": "No re-describir lo que el anchor ya fija; SW30 R16"},
    {"codigo": "E6.3", "etapa": "E6", "nombre": "Bloques de continuidad", "autoridad": "video/universal-rules U7 y U13", "entra": "Secuencia de clips", "sale": "Identidad completa repetida en cada prompt, rol de cada referencia", "gate": "Cada asset con rol explícito"},
    {"codigo": "E6.4", "etapa": "E6", "nombre": "Dramaturgy check + auditoría de 3 detalles", "autoridad": "video/dramaturgy.md §15, universal-rules U12", "entra": "Cada prompt", "sale": "Seis respuestas sí, tres detalles por shot", "gate": "Si una es no, no se entrega"},
    {"codigo": "E6.5", "etapa": "E6", "nombre": "Linter determinista", "autoridad": "aurora-prompt-linter", "entra": "Prompt + refs.yaml + caso", "sale": "PASS o FAIL con violaciones", "gate": "prompt_linter.py exit 0"},
    {"codigo": "E6.6", "etapa": "E6", "nombre": "Micro-gate y entrega", "autoridad": "produccion-visual-sw30", "entra": "Prompt que pasó todo", "sale": "Cabecera SKILL / RIESGOS / TÉCNICA + prompt", "gate": "Entrega sin cabecera es inválida"},
    {"codigo": "E7.1", "etapa": "E7", "nombre": "Ensamblar el árbol del paquete", "autoridad": "production-package.md §1", "entra": "Todos los artefactos", "sale": "production-package/ 00 a 07", "gate": "Estructura exacta de la spec"},
    {"codigo": "E7.2", "etapa": "E7", "nombre": "Preview de cliente", "autoridad": "storyboard-html-preview", "entra": "shots, overlays, snapshot, frames, critiques", "sale": "preview.html single-file offline", "gate": "shots-to-html.py --selftest; sin CDN"},
    {"codigo": "E7.3", "etapa": "E7", "nombre": "Checklist de post-producción", "autoridad": "production-package.md §3", "entra": "Dirección de audio, grade, plataforma", "sale": "post-production.md con specs concretos", "gate": "Sin placeholders"},
    {"codigo": "E7.4", "etapa": "E7", "nombre": "Trazabilidad y créditos", "autoridad": "production-package.md §2 y §4", "entra": "Paquete completo", "sale": "3 prompts al azar trazan a su concepto; atribución smixs", "gate": "Si un eslabón falta, no se entrega"},
    {"codigo": "POST.1", "etapa": "POST", "nombre": "Generar clips y verificar contra shot card", "autoridad": "Checklist §3", "entra": "Prompts de E6", "sale": "Clips aprobados", "gate": "Encuadre, acción, luz"},
    {"codigo": "POST.2", "etapa": "POST", "nombre": "Ensamblar, grade, audio, overlays, export", "autoridad": "Checklist §3", "entra": "Clips + text-overlays.json + brand-lock", "sale": "Master exportado por plataforma", "gate": "QA final contra los 5 anclas"},
]

# Tareas sueltas de app/build_app.py (briefs que no son un video completo). Se leen del
# propio build_app.py cuando se puede importar; esto es el respaldo con los mismos nombres.
SUELTAS_RESPALDO: dict[str, dict] = {
    "PROMPT_IMAGEN": {"nombre": "Prompt de imagen", "casos": ["T1", "T2", "T3", "T4", "T5", "PROD", "GRAF", "MULTI"]},
    "PROMPT_VIDEO": {"nombre": "Prompt de video", "casos": ["CLIP"]},
    "REF": {"nombre": "Referencia a prompt", "casos": ["REF"]},
    "ESTRATEGIA": {"nombre": "Estrategia", "casos": []},
    "DOC": {"nombre": "Documento", "casos": []},
}

# Facetas de la tarjeta de ancla. Las descripciones alimentan los prototipos.
FACETAS: dict[str, dict] = {
    "d1": {"nombre": "rol", "cerrada": 1, "valores": {
        "character_ref": "referencia de personaje: maestro de identidad, rostro o cuerpo canonizado",
        "environment_ref": "referencia de entorno: placa de lugar vacía, locación canonizada",
        "keyframe_ff": "keyframe de inicio (first frame) de un clip: la imagen de arranque",
        "keyframe_lf": "keyframe final (last frame, tail) de un clip: la imagen de llegada",
        "hero_macro": "macro del objeto héroe: producto, detalle, textura en primer plano",
        "motif": "motivo visual recurrente: el objeto o forma que se repite en la pieza",
    }},
    "d2": {"nombre": "accion", "cerrada": 1, "valores": {
        "A_pose": "pose estática: retrato, postura sostenida, sujeto quieto",
        "B_pico": "instante pico congelado: apex del salto, momento de impacto, gesto en su punto máximo",
        "C_movimiento": "movimiento continuo leído en un still: velocidad, barrido, estela, desplazamiento",
        "D_luz_clima": "la acción es la luz o el clima: rayo, lluvia, amanecer, niebla, polvo",
        "E_interaccion": "interacción entre sujetos: contacto, mirada, gesto compartido, high five",
        "F_multitud": "multitud como acción: masa que se mueve, público, estadio, marea de gente",
    }},
    "d3": {"nombre": "sujetos", "cerrada": 1, "valores": {
        "0": "sin sujeto humano: objeto, lugar, producto, vehículo sin piloto visible",
        "1": "un solo sujeto en cuadro",
        "2_sin_contacto": "dos sujetos sin contacto físico: conversación, distancia, mirada",
        "2_contacto": "dos sujetos con contacto físico: abrazo, high five, apretón, choque",
        "ensamble": "grupo con identidad, cada miembro reconocible: equipo, banda, familia",
        "multitud": "multitud anónima, sin rostros legibles: estadio, público, calle llena",
    }},
    "d4": {"nombre": "tratamiento", "cerrada": 1, "valores": {
        "documental": "tratamiento documental: luz disponible, candid, sin puesta en escena visible",
        "narrativo": "tratamiento narrativo o cinematográfico: escena de ficción, drama, cortometraje",
        "comercial": "tratamiento comercial: spot, producto, estética publicitaria pulida",
        "race": "tratamiento de carreras y velocidad: drift, drag, chase, montaje cinético sin rostros",
        "ugc": "tratamiento UGC: contenido de usuario, celular en mano, TikTok, reel orgánico",
        "animacion": "tratamiento de animación: 2D, 3D, stop motion, ilustración en movimiento",
    }},
    "d5": {"nombre": "paleta", "cerrada": 0, "valores": {}},
    "d6": {"nombre": "encuadre_lente", "cerrada": 0, "valores": {}},
    "d7": {"nombre": "luz", "cerrada": 0, "valores": {}},
    "d8": {"nombre": "modelo", "cerrada": 1, "valores": {
        "nano-banana-2": "Nano Banana 2 (NB2): imagen Google, grounding, ratios extremos, barato",
        "nano-banana-pro": "Nano Banana Pro (NBP): imagen Google, escenas complejas con física y hasta 14 referencias",
        "gpt-image-2": "GPT Image 2: imagen OpenAI, texto fino, UI, edición con preservación de identidad",
        "midjourney": "Midjourney: imagen estética, parámetros --ar --stylize, sin sintaxis de referencia por texto",
        "flux": "Flux: imagen open weights, prompts descriptivos largos",
        "ideogram": "Ideogram: imagen con tipografía y texto renderizado",
        "seedream": "Seedream: imagen ByteDance",
        "kling": "Kling: video Kuaishou, física realista, multi-shot en 3.0, negative prompt en campo aparte",
        "veo": "Veo: video Google, audio nativo, diálogo y lip-sync",
        "seedance": "Seedance: video ByteDance, multi-shot nativo, storytelling en un prompt",
        "hailuo": "Hailuo: video MiniMax, movimiento de cámara por comandos",
        "sora": "Sora: video OpenAI",
        "agnostico": "prompt agnóstico de modelo: artefacto intermedio, storyboard, shot card",
    }},
    "d9": {"nombre": "angulo", "cerrada": 1, "valores": {
        "eye-level": "ángulo a nivel de ojos: cámara a la altura del sujeto",
        "high": "ángulo alto: cámara por encima del sujeto mirando hacia abajo",
        "low": "ángulo bajo: cámara por debajo del sujeto mirando hacia arriba, a ras de suelo",
        "overhead": "cenital: cámara vertical desde arriba, top-down",
        "dutch": "ángulo holandés: horizonte inclinado, cámara ladeada",
        "pov": "punto de vista subjetivo: la cámara es los ojos del sujeto, primera persona",
    }},
}

# --------------------------------------------------------------------------- derivaciones por ruta
# Todas las tablas de aquí abajo se leen "el primer patrón que hace match manda" salvo
# las de facetas, donde se acumulan todos los que hacen match.

# medio por skill/ruta
MEDIO: list[tuple[str, str]] = [
    ("video/**", "VIDEO"),
    ("image/**", "IMAGEN"),
    ("visual-prompt-forge/adapters/kling.md", "VIDEO"),
    ("visual-prompt-forge/adapters/veo.md", "VIDEO"),
    ("visual-prompt-forge/adapters/seedance.md", "VIDEO"),
    ("visual-prompt-forge/adapters/hailuo.md", "VIDEO"),
    ("visual-prompt-forge/adapters/*.md", "IMAGEN"),
    ("visual-prompt-forge/**", "AMBOS"),
    ("visual-asset-critic/**", "IMAGEN"),
    ("ai-video-storyboard/**", "VIDEO"),
    ("storyboard-architect/**", "AMBOS"),
    ("screenwriter/**", "TEXTO"),
    ("brand-lock-extractor/**", "PROCESO"),
    ("storyboard-html-preview/**", "PROCESO"),
    ("ai-production-director/**", "PROCESO"),
    ("produccion-visual-sw30/**", "AMBOS"),
    ("aurora-prompt-linter/**", "AMBOS"),
    ("visual-media/**", "AMBOS"),
    ("regimenes/**", "IMAGEN"),
]

FASE: list[tuple[str, str]] = [
    ("ai-production-director/references/creative-strategy.md", "PLANEACION"),
    ("ai-production-director/references/production-package.md", "ENTREGA"),
    ("ai-production-director/**", "TRANSVERSAL"),
    ("screenwriter/**", "PLANEACION"),
    ("video/references/dramaturgy.md", "PLANEACION"),
    ("video/references/role-modes.md", "PLANEACION"),
    ("video/references/patterns-and-genres.md", "PLANEACION"),
    ("video/**", "REDACCION"),
    ("image/**", "REDACCION"),
    ("visual-prompt-forge/**", "REDACCION"),
    ("regimenes/**", "REDACCION"),
    ("aurora-prompt-linter/**", "REDACCION"),
    ("storyboard-architect/**", "PLANEACION"),
    ("ai-video-storyboard/**", "PLANEACION"),
    ("brand-lock-extractor/**", "PLANEACION"),
    ("visual-asset-critic/**", "QA_RENDER"),
    ("storyboard-html-preview/**", "ENTREGA"),
    ("produccion-visual-sw30/**", "TRANSVERSAL"),
    ("visual-media/**", "TRANSVERSAL"),
]

# gate mecánico que verifica la regla (ruta + sección/texto). Sin gate => MANUAL.
GATES: list[tuple[str, re.Pattern | None, str]] = [
    ("image/references/golden-rules.md", None, "gate_image"),
    ("video/references/dramaturgy.md", re.compile(r"banned", re.I), "gate_dramaturgy"),
    ("image/references/gpt-image.md", re.compile(r"stunning|masterpiece|gorgeous", re.I), "gate_dramaturgy"),
    ("ai-production-director/SKILL.md", re.compile(r"6\.1 Vocabulario", re.I), "gate_dramaturgy"),
    ("storyboard-architect/references/timing-rules.md", None, "validate_shots"),
    ("storyboard-architect/SKILL.md", re.compile(r"shots\.json|validate", re.I), "validate_shots"),
    ("aurora-prompt-linter/**", None, "aurora"),
    ("produccion-visual-sw30/SKILL.md", re.compile(r"MICRO-GATE", re.I), "gate_microgate"),
]

# caso por ruta (origen 'archivo'), sólo para reglas que aún no tienen caso.
ARCHIVO_CASO: list[tuple[str, list[str]]] = [
    ("image/references/patterns/ecommerce.md", ["PROD"]),
    ("image/references/patterns/food-beverage.md", ["PROD"]),
    ("image/references/patterns/fashion-editorial.md", ["T2", "T3"]),
    ("image/references/patterns/portrait-cinema.md", ["T1", "T3"]),
    ("image/references/patterns/character-design.md", ["T1", "T2"]),
    ("image/references/patterns/poster-illustration.md", ["GRAF"]),
    ("image/references/patterns/ui-social.md", ["GRAF"]),
    ("image/references/slides.md", ["GRAF"]),
    ("image/references/text-rendering.md", ["GRAF"]),
    ("image/references/multi-panel.md", ["MULTI"]),
    ("image/references/storyboards.md", ["MULTI", "SHOT"]),
    ("image/references/vision-decomposer.md", ["REF"]),
    ("image/references/editing.md", ["T5"]),
    ("image/references/characters.md", ["T1", "T2"]),
    ("image/**", ["T1", "T2", "T3", "T4", "T5", "PROD", "GRAF", "MULTI"]),
    ("video/references/animatic-keyframes.md", ["SHOT", "T3", "OBJ", "LUGAR"]),
    ("video/references/race-and-speed.md", ["SHOT", "OBJ", "CLIP"]),
    ("video/references/dramaturgy.md", ["CLIP", "SHOT", "GUION"]),
    ("video/**", ["CLIP"]),
    ("visual-prompt-forge/adapters/kling.md", ["CLIP"]),
    ("visual-prompt-forge/adapters/veo.md", ["CLIP"]),
    ("visual-prompt-forge/adapters/seedance.md", ["CLIP"]),
    ("visual-prompt-forge/adapters/hailuo.md", ["CLIP"]),
    ("visual-prompt-forge/adapters/*.md", ["T1", "T2", "T3", "T4", "PROD", "GRAF"]),
    ("visual-prompt-forge/**", ["SHOT", "CLIP", "T3"]),
    ("visual-asset-critic/**", ["QA"]),
    ("storyboard-architect/**", ["SHOT"]),
    ("ai-video-storyboard/**", ["SHOT", "CLIP"]),
    ("screenwriter/**", ["GUION"]),
    ("brand-lock-extractor/**", ["MARCA"]),
    ("storyboard-html-preview/**", ["ENTREGA"]),
    ("ai-production-director/references/creative-strategy.md", ["GUION"]),
    ("ai-production-director/references/production-package.md", ["ENTREGA"]),
    ("aurora-prompt-linter/**", ["CLIP", "T1", "T2", "T3", "T4", "T5", "PROD", "GRAF"]),
    ("visual-media/**", ["NINGUNO"]),
    # regímenes: por nombre de archivo, cuando el frontmatter no declara `casos:`
    ("regimenes/*multitud*", ["MULTITUD"]),
    ("regimenes/*ensamble*", ["ENSAMBLE"]),
    ("regimenes/*grupo*", ["ENSAMBLE"]),
    ("regimenes/*vehiculo*", ["OBJ"]),
    ("regimenes/*objeto*", ["OBJ"]),
    ("regimenes/*balistico*", ["OBJ"]),
    ("regimenes/*luz*", ["LUGAR"]),
    ("regimenes/*clima*", ["LUGAR"]),
    ("regimenes/*agua*", ["LUGAR", "T3"]),
    ("regimenes/*subacuatico*", ["LUGAR", "T3"]),
    ("regimenes/*cuerpo*", ["T3", "T2"]),
    ("regimenes/**", ["T3"]),
]

# tarea por ruta (origen 'derivada_archivo'): columna "Skill con autoridad" de la tabla
# de subprocesos, invertida. Todos los que hagan match acumulan.
ARCHIVO_TAREA: list[tuple[str, list[str]]] = [
    ("brand-lock-extractor/**", ["E0.1", "E0.2", "E0.3", "E0.4"]),
    ("ai-production-director/references/creative-strategy.md", ["E1.1", "E1.2", "E1.3", "E1.4", "E1.5", "ESTRATEGIA"]),
    ("ai-production-director/references/production-package.md", ["E7.1", "E7.3", "E7.4"]),
    ("ai-production-director/SKILL.md", ["E5.1"]),
    ("screenwriter/**", ["E2.1", "E2.2", "E2.3", "E2.4"]),
    ("video/references/dramaturgy.md", ["E3.1", "E3.2", "E3.3", "E3.4", "E3.5", "E6.4"]),
    ("video/references/universal-rules.md", ["E3.3", "E6.3", "E6.4"]),
    ("video/references/camera-lighting-vocabulary.md", ["E3.3", "E3.4"]),
    ("video/references/role-modes.md", ["E3.2", "E4.3"]),
    ("video/references/patterns-and-genres.md", ["E3.4", "E1.5"]),
    ("video/references/kling.md", ["E6.2"]),
    ("video/references/veo.md", ["E6.2"]),
    ("video/references/seedance.md", ["E6.2"]),
    ("video/references/seedance-25.md", ["E6.2"]),
    ("video/references/fixes-and-skeletons.md", ["E6.2", "E6.4"]),
    ("video/references/animatic-keyframes.md", ["E5.2", "E5.4", "E4.3"]),
    ("video/references/race-and-speed.md", ["E5.2", "E5.4", "E6.2"]),
    ("video/SKILL.md", ["E6.2", "E6.4", "PROMPT_VIDEO"]),
    ("storyboard-architect/references/beat-frameworks.md", ["E4.1"]),
    ("storyboard-architect/references/timing-rules.md", ["E4.2"]),
    ("storyboard-architect/references/shot-grammar.md", ["E4.3"]),
    ("storyboard-architect/references/on-screen-text.md", ["E4.4"]),
    ("storyboard-architect/brand-packs/**", ["E4.1", "E0.4"]),
    ("storyboard-architect/tools/**", ["E4.6"]),
    ("storyboard-architect/examples/**", ["E4.6"]),
    ("storyboard-architect/SKILL.md", ["E4.1", "E4.2", "E4.3", "E4.4", "E4.5", "E4.6"]),
    ("ai-video-storyboard/**", ["E4.3"]),
    ("image/references/editing.md", ["E5.5"]),
    ("image/references/characters.md", ["E5.2", "E5.3"]),
    ("image/references/patterns/**", ["E5.2", "E5.4"]),
    ("image/references/vision-decomposer.md", ["REF"]),
    ("image/references/multi-panel.md", ["E5.2"]),
    ("image/references/storyboards.md", ["E5.4", "E4.3"]),
    ("image/references/text-rendering.md", ["E5.2"]),
    ("image/references/slides.md", ["E5.2"]),
    ("image/**", ["E5.2", "E5.3", "E5.4", "E5.5", "PROMPT_IMAGEN"]),
    ("visual-prompt-forge/adapters/kling.md", ["E6.1"]),
    ("visual-prompt-forge/adapters/veo.md", ["E6.1"]),
    ("visual-prompt-forge/adapters/seedance.md", ["E6.1"]),
    ("visual-prompt-forge/adapters/hailuo.md", ["E6.1"]),
    ("visual-prompt-forge/adapters/*.md", ["E5.6"]),
    ("visual-prompt-forge/references/consistency-locks.md", ["E5.6", "E6.3"]),
    ("visual-prompt-forge/references/failure-modes.md", ["E5.8"]),
    ("visual-prompt-forge/**", ["E5.6", "E5.8", "E6.1"]),
    ("visual-asset-critic/**", ["E5.7"]),
    ("produccion-visual-sw30/**", ["E5.2", "E5.3", "E5.4", "E5.5", "E5.9", "E6.2", "E6.6"]),
    ("aurora-prompt-linter/**", ["E6.5"]),
    ("storyboard-html-preview/**", ["E7.2"]),
    ("visual-media/**", ["DOC"]),  # director §7: sólo animación y material didáctico
    ("regimenes/**", ["E5.2", "E5.4"]),
]

# refinamiento caso -> subprocesos (origen 'derivada_caso'). Lo que no esté aquí toma
# todos los subprocesos de la etapa que build_app.py asigna al caso.
CASO_SUBPROCESO: dict[str, list[str]] = {
    "T1": ["E5.2"], "T2": ["E5.3"], "T3": ["E5.4"], "T4": ["E5.4"], "T5": ["E5.5"],
    "PROD": ["E5.2", "E5.4"], "GRAF": ["E5.2"], "MULTI": ["E5.2"], "REF": ["REF"],
    "LUGAR": ["E5.2"], "EDIF": ["E5.2"], "ANIMAL": ["E5.2", "E5.3"], "OBJ": ["E5.2", "E5.4"],
    "MULTITUD": ["E5.4"], "ENSAMBLE": ["E5.4"], "QA": ["E5.7", "E5.8"],
    "CLIP": ["E6.1", "E6.2", "E6.3", "E6.4", "E6.5", "E6.6"],
}
# etapas para los casos nuevos, que build_app.py aún no conoce
ETAPA_EXTRA: dict[str, list[str]] = {
    "LUGAR": ["E5", "PROMPT_IMAGEN"], "EDIF": ["E5", "PROMPT_IMAGEN"], "ANIMAL": ["E5", "PROMPT_IMAGEN"],
    "OBJ": ["E5", "PROMPT_IMAGEN"], "MULTITUD": ["E5", "PROMPT_IMAGEN"], "ENSAMBLE": ["E5", "PROMPT_IMAGEN"],
}

# facetas por ruta (origen 'archivo'); acumulan.
ARCHIVO_FACETA: list[tuple[str, str, list[str]]] = [
    ("video/references/kling.md", "d8", ["kling"]),
    ("video/references/veo.md", "d8", ["veo"]),
    ("video/references/seedance.md", "d8", ["seedance"]),
    ("video/references/seedance-25.md", "d8", ["seedance"]),
    ("visual-prompt-forge/adapters/kling.md", "d8", ["kling"]),
    ("visual-prompt-forge/adapters/veo.md", "d8", ["veo"]),
    ("visual-prompt-forge/adapters/seedance.md", "d8", ["seedance"]),
    ("visual-prompt-forge/adapters/hailuo.md", "d8", ["hailuo"]),
    ("visual-prompt-forge/adapters/midjourney.md", "d8", ["midjourney"]),
    ("visual-prompt-forge/adapters/flux.md", "d8", ["flux"]),
    ("visual-prompt-forge/adapters/ideogram.md", "d8", ["ideogram"]),
    ("visual-prompt-forge/adapters/seedream.md", "d8", ["seedream"]),
    ("visual-prompt-forge/adapters/gpt-image.md", "d8", ["gpt-image-2"]),
    ("visual-prompt-forge/adapters/nano-banana.md", "d8", ["nano-banana-2", "nano-banana-pro"]),
    ("image/references/gpt-image.md", "d8", ["gpt-image-2"]),
    ("image/references/nano-banana.md", "d8", ["nano-banana-2", "nano-banana-pro"]),
    ("video/references/race-and-speed.md", "d2", ["C_movimiento"]),
    ("video/references/race-and-speed.md", "d4", ["race"]),
    ("video/references/race-and-speed.md", "d3", ["0"]),
    ("video/references/animatic-keyframes.md", "d1", ["keyframe_ff", "keyframe_lf"]),
    ("image/references/characters.md", "d1", ["character_ref"]),
    ("image/references/patterns/character-design.md", "d1", ["character_ref"]),
    ("image/references/patterns/portrait-cinema.md", "d2", ["A_pose"]),
    ("image/references/patterns/portrait-cinema.md", "d4", ["narrativo"]),
    ("image/references/patterns/ecommerce.md", "d1", ["hero_macro"]),
    ("image/references/patterns/ecommerce.md", "d4", ["comercial"]),
    ("image/references/patterns/food-beverage.md", "d1", ["hero_macro"]),
    ("image/references/patterns/food-beverage.md", "d4", ["comercial"]),
    ("image/references/patterns/fashion-editorial.md", "d4", ["comercial"]),
    ("image/references/patterns/fashion-editorial.md", "d2", ["A_pose"]),
    ("visual-prompt-forge/references/consistency-locks.md", "d1", ["character_ref", "environment_ref"]),
    ("ai-video-storyboard/**", "d4", ["ugc"]),
    ("visual-media/**", "d4", ["animacion"]),
]

# facetas por regex sobre el texto (origen 'regex'); acumulan. (dimension, valor, patron)
REGEX_FACETA: list[tuple[str, str, re.Pattern]] = [
    ("d8", "kling", re.compile(r"\bkling\b", re.I)),
    ("d8", "veo", re.compile(r"\bveo\b", re.I)),
    ("d8", "seedance", re.compile(r"\bseedance\b", re.I)),
    ("d8", "hailuo", re.compile(r"\bhailuo\b|\bminimax\b", re.I)),
    ("d8", "sora", re.compile(r"\bsora\b", re.I)),
    ("d8", "midjourney", re.compile(r"\bmidjourney\b|\bMJ\b", re.I)),
    ("d8", "flux", re.compile(r"\bflux\b", re.I)),
    ("d8", "ideogram", re.compile(r"\bideogram\b", re.I)),
    ("d8", "seedream", re.compile(r"\bseedream\b", re.I)),
    ("d8", "nano-banana-pro", re.compile(r"nano[ -]?banana[ -]?pro|\bNBP\b", re.I)),
    ("d8", "nano-banana-2", re.compile(r"nano[ -]?banana[ -]?2\b|\bNB2\b", re.I)),
    ("d8", "gpt-image-2", re.compile(r"gpt[ -]?image", re.I)),
    ("d9", "eye-level", re.compile(r"eye[ -]level|a nivel de ojos|altura de los ojos", re.I)),
    ("d9", "high", re.compile(r"high[ -]angle|from above|bird'?s[ -]eye|picado|ángulo alto|angulo alto|desde arriba", re.I)),
    ("d9", "low", re.compile(r"low[ -]angle|from below|worm'?s[ -]eye|contrapicado|ángulo bajo|angulo bajo|desde abajo|a ras de", re.I)),
    ("d9", "overhead", re.compile(r"\boverhead\b|top[ -]down|cenital", re.I)),
    ("d9", "dutch", re.compile(r"dutch (?:angle|tilt)|canted|holand[ée]s", re.I)),
    ("d9", "pov", re.compile(r"\bPOV\b|point[ -]of[ -]view|first[ -]person|primera persona|subjetiv[oa]", re.I)),
    ("d2", "A_pose", re.compile(r"\bposes?\b|\bposture\b|\bportrait\b|\bretrato\b|\bpostura\b", re.I)),
    ("d2", "B_pico", re.compile(r"peak (?:action|moment)|mid-?air|frozen mid|\bapex\b|moment of impact|instante|pico de la acci[oó]n|congelad[oa]", re.I)),
    ("d2", "C_movimiento", re.compile(r"motion blur|\bvelocity\b|\bspeed\b|\bvelocidad\b|\bchase\b|\bdrift\b|\bracing\b|\bestela\b|barrid[oa]", re.I)),
    ("d2", "D_luz_clima", re.compile(r"\brain\b|\bfog\b|\bsnow\b|\blightning\b|\bstorm\b|\bweather\b|\blluvia\b|\bniebla\b|\bnieve\b|\brayo\b|\btormenta\b|\bclima\b|golden hour|amanecer|atardecer", re.I)),
    ("d2", "E_interaccion", re.compile(r"high[ -]five|handshake|\bhug\b|eye contact|\binteraction\b|\babrazo\b|contacto|interacci[oó]n", re.I)),
    ("d2", "F_multitud", re.compile(r"\bcrowd\b|\bstadium\b|\baudience\b|\bmultitud\b|\bestadio\b|\bp[úu]blico\b", re.I)),
    ("d3", "0", re.compile(r"\bno people\b|\bfaceless\b|faces (?:are )?banned|\bempty (?:plate|room|street|set)\b|sin personas?\b|sin rostro", re.I)),
    ("d3", "1", re.compile(r"\bone person\b|\bsingle (?:subject|person|character)\b|\bsolo subject\b|\buna persona\b|\bun solo sujeto\b|\b1 persona\b", re.I)),
    ("d3", "2_sin_contacto", re.compile(r"\btwo people\b|\b2 people\b|\bdos personas\b|\b2 personas\b|\bpair\b|\bcouple\b|two-?shot", re.I)),
    ("d3", "2_contacto", re.compile(r"high[ -]five|handshake|\bhug\b|\bembrace\b|\babrazo\b|contacto entre personas|two people touching", re.I)),
    ("d3", "ensamble", re.compile(r"\bgroup shot\b|\bteam\b|\bband\b|\bfamily\b|\bensemble\b|\bensamble\b|\bequipo\b|\bfamilia\b|\bgrupo\b", re.I)),
    ("d3", "multitud", re.compile(r"\bcrowd\b|\bmultitud\b|\bestadio\b|\bstadium\b", re.I)),
    ("d4", "documental", re.compile(r"\bdocumentary\b|\bdocumental\b|\bcandid\b|available light", re.I)),
    ("d4", "narrativo", re.compile(r"\bnarrative\b|\bscreenplay\b|short film|\bdramatic\b|\bnarrativ[oa]\b|\bcortometraje\b", re.I)),
    ("d4", "comercial", re.compile(r"\bcommercial\b|\badvert(?:ising|isement)?\b|product shot|\bcomercial\b|\bspot\b|\bpublicitari[oa]\b", re.I)),
    ("d4", "race", re.compile(r"\brace\b|\bracing\b|\bdrift\b|drag[ -]strip|\bcarreras?\b", re.I)),
    ("d4", "ugc", re.compile(r"\bUGC\b|user[ -]generated|\btiktok\b|\breels?\b|\bshorts\b", re.I)),
    ("d4", "animacion", re.compile(r"\banimation\b|\banimated\b|\bcartoon\b|stop[ -]motion|\banimaci[oó]n\b", re.I)),
    ("d1", "character_ref", re.compile(r"character (?:ref|reference|sheet)|identity lock|\bmaestro de (?:rostro|cuerpo|persona)\b|face master|reference face", re.I)),
    ("d1", "environment_ref", re.compile(r"environment (?:ref|reference|plate)|empty plate|placa (?:de lugar|vac[ií]a)|location plate", re.I)),
    ("d1", "keyframe_ff", re.compile(r"first frame|start frame|start image|keyframe de inicio|frame inicial|\bFF\b", re.I)),
    ("d1", "keyframe_lf", re.compile(r"last frame|end frame|tail image|tail frame|final frame|keyframe final|\bLF\b", re.I)),
    ("d1", "hero_macro", re.compile(r"hero (?:shot|product|macro)|\bmacro\b|producto h[ée]roe", re.I)),
    ("d1", "motif", re.compile(r"\bmotif\b|\bmotivo visual\b|\bleitmotiv\b", re.I)),
    # dimensiones abiertas: el valor es el token normalizado que aparece
    ("d6", "wide-shot", re.compile(r"\bwide shot\b|\bplano general\b|\bWS\b", re.I)),
    ("d6", "medium-shot", re.compile(r"\bmedium shot\b|\bplano medio\b|\bMS\b", re.I)),
    ("d6", "close-up", re.compile(r"\bclose[ -]up\b|\bprimer plano\b|\bCU\b", re.I)),
    ("d6", "extreme-close-up", re.compile(r"extreme close[ -]up|\bECU\b|primer[ií]simo", re.I)),
    ("d6", "full-shot", re.compile(r"\bfull shot\b|\bfull[ -]body\b|\bcuerpo entero\b", re.I)),
    ("d6", "over-the-shoulder", re.compile(r"over[ -]the[ -]shoulder|\bOTS\b", re.I)),
    ("d6", "two-shot", re.compile(r"\btwo[ -]shot\b", re.I)),
    ("d6", "macro", re.compile(r"\bmacro (?:lens|shot)\b", re.I)),
    ("d6", "telephoto", re.compile(r"\btelephoto\b|\bteleobjetivo\b|\b(?:85|100|135|200)\s?mm\b", re.I)),
    ("d6", "wide-lens", re.compile(r"\bwide[ -]angle\b|\bgran angular\b|\b(?:14|16|18|20|24)\s?mm\b", re.I)),
    ("d6", "normal-lens", re.compile(r"\b(?:35|50)\s?mm\b", re.I)),
    ("d6", "anamorphic", re.compile(r"\banamorphic\b|\banam[oó]rfic[oa]\b", re.I)),
    ("d7", "golden-hour", re.compile(r"golden hour|hora dorada", re.I)),
    ("d7", "blue-hour", re.compile(r"blue hour|hora azul", re.I)),
    ("d7", "backlight", re.compile(r"\bbacklight|\bbacklit\b|\bcontraluz\b", re.I)),
    ("d7", "rim-light", re.compile(r"\brim[ -]light|\brim\b", re.I)),
    ("d7", "soft-light", re.compile(r"\bsoft light|\bdiffused light|\bluz suave\b", re.I)),
    ("d7", "hard-light", re.compile(r"\bhard light|\bharsh light|\bluz dura\b", re.I)),
    ("d7", "overcast", re.compile(r"\bovercast\b|\bnublado\b", re.I)),
    ("d7", "neon", re.compile(r"\bneon\b|\bne[oó]n\b", re.I)),
    ("d7", "practical", re.compile(r"\bpracticals?\b|practical light", re.I)),
    ("d7", "window-light", re.compile(r"window light|luz de ventana", re.I)),
    ("d7", "low-key", re.compile(r"\blow[ -]key\b|\bchiaroscuro\b", re.I)),
    ("d7", "high-key", re.compile(r"\bhigh[ -]key\b", re.I)),
    ("d7", "silhouette", re.compile(r"\bsilhouette\b|\bsilueta\b", re.I)),
    ("d7", "three-point", re.compile(r"three[ -]point", re.I)),
    ("d5", "teal-orange", re.compile(r"teal (?:and|&) orange", re.I)),
    ("d5", "monochrome", re.compile(r"\bmonochrom\w*|\bblack and white\b|\bB&W\b|\bblanco y negro\b", re.I)),
    ("d5", "desaturated", re.compile(r"\bdesaturated\b|\bdesaturad[oa]\b", re.I)),
    ("d5", "pastel", re.compile(r"\bpastel\b", re.I)),
    ("d5", "warm", re.compile(r"\bwarm (?:tones?|palette|grade|amber)\b|\btonos c[áa]lidos\b", re.I)),
    ("d5", "cold", re.compile(r"\bcold (?:tones?|palette|grade)\b|\bcool tones?\b|\bsteel blue\b|\bcyan\b|\btonos fr[íi]os\b|paleta fr[íi]a", re.I)),
    ("d5", "hex", re.compile(r"#[0-9a-fA-F]{6}\b")),
]


# --------------------------------------------------------------------------- utilidades

def now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def sha1(texto: str) -> str:
    return hashlib.sha1(texto.encode("utf-8")).hexdigest()


def connect(path: str) -> sqlite3.Connection:
    con = sqlite3.connect(path)
    con.execute("PRAGMA foreign_keys = ON")
    return con


def ruta(skill: str, archivo: str) -> str:
    return f"{skill}/{archivo}"


def primero(tabla: list[tuple[str, str]], skill: str, archivo: str, defecto: str) -> str:
    r = ruta(skill, archivo)
    for patron, valor in tabla:
        if fnmatch.fnmatch(r, patron):
            return valor
    return defecto


def acumular(tabla, skill: str, archivo: str) -> list:
    r = ruta(skill, archivo)
    out: list = []
    for fila in tabla:
        if fnmatch.fnmatch(r, fila[0]):
            out.append(fila[1:])
    return out


def cargar_build_app() -> dict:
    """ETAPAS de app/build_app.py sin ejecutar su main. Respaldo si no se puede importar."""
    try:
        spec = importlib.util.spec_from_file_location("build_app", BUILD_APP)
        mod = importlib.util.module_from_spec(spec)  # type: ignore[arg-type]
        spec.loader.exec_module(mod)  # type: ignore[union-attr]
        return {"ETAPAS": mod.ETAPAS, "TRACKS": getattr(mod, "TRACKS", {}), "fuente": str(BUILD_APP)}
    except Exception as exc:  # noqa: BLE001
        etapas = [{"clave": k, "nombre": v["nombre"], "casos": v["casos"], "tracks": [],
                   "cond": "", "entra": "", "sale": "", "gate": ""} for k, v in SUELTAS_RESPALDO.items()]
        return {"ETAPAS": etapas, "TRACKS": {}, "fuente": f"respaldo ({exc})"}


def caso_a_etapas(etapas: list[dict]) -> dict[str, list[str]]:
    mapa: dict[str, list[str]] = {}
    for e in etapas:
        for c in e.get("casos", []):
            mapa.setdefault(c, []).append(e["clave"])
    for c, claves in ETAPA_EXTRA.items():
        for k in claves:
            mapa.setdefault(c, [])
            if k not in mapa[c]:
                mapa[c].append(k)
    return mapa


def capas_del_director() -> dict[str, list[int]]:
    """Parsea la tabla '## 1. Mapa de sub-skills' del director: skill -> capas numéricas."""
    capas: dict[str, list[int]] = {}
    roots = rr.skill_roots("ai-production-director")
    if not roots:
        return capas
    texto = (roots[0][0] / "SKILL.md").read_text(encoding="utf-8")
    en_tabla = False
    for linea in texto.splitlines():
        if linea.startswith("## 1."):
            en_tabla = True
            continue
        if en_tabla and linea.startswith("## "):
            break
        if not en_tabla or not linea.startswith("|"):
            continue
        celdas = [c.strip() for c in linea.strip().strip("|").split("|")]
        if len(celdas) < 2 or celdas[0].startswith("---") or celdas[0].lower().startswith("sub-skill"):
            continue
        m = re.search(r"`([\w-]+)`", celdas[0])
        nombre = m.group(1) if m else ("ai-production-director" if "este skill" in celdas[0] else "")
        if not nombre:
            continue
        numeros = [int(n) for n in re.findall(r"\b(\d)\b", celdas[1])]
        if len(numeros) == 2 and re.search(r"\d\s*[–-]\s*\d", celdas[1]):
            numeros = list(range(numeros[0], numeros[1] + 1))
        capas.setdefault(nombre, [])
        for n in numeros:
            if n not in capas[nombre]:
                capas[nombre].append(n)
    return capas


def parse_flujo_html(path: Path) -> list[dict]:
    """Regenera TAREAS desde flujo-spot.html (sección 'Subprocesos por etapa')."""
    import html as htmlmod

    t = path.read_text(encoding="utf-8")
    i = t.find("<h2>Subprocesos por etapa</h2>")
    j = t.find("</table>", i)
    filas = re.findall(r"<tr[^>]*>(.*?)</tr>", t[i:j], re.S)
    tareas: list[dict] = []
    etapa = ""
    for fila in filas:
        celdas = [htmlmod.unescape(re.sub(r"<[^>]+>", "", c)).strip()
                  for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", fila, re.S)]
        if len(celdas) == 2:
            etapa = celdas[0]
        elif len(celdas) == 6 and celdas[0] != "Código":
            tareas.append({"codigo": celdas[0], "etapa": etapa, "nombre": celdas[1],
                           "autoridad": celdas[2], "entra": celdas[3], "sale": celdas[4],
                           "gate": celdas[5]})
    return tareas


def prioridad_de(regla: dict, gate: str | None) -> int:
    estado = regla["estado"]
    if estado == "CANONICA":
        return 1
    if estado == "REFUTADA":
        return 5
    if estado in ("A_PRUEBA", "CAMPO"):
        return 4
    p = {"PROHIBICION": 2, "OBLIGACION": 2, "RECOMENDACION": 3, "EJEMPLO": 4, "ESTRUCTURA": 4}[regla["fuerza"]]
    if regla["fuerza"] == "PROHIBICION" and gate == "gate_dramaturgy":
        p = 1  # §6.1 del director: el vocabulario prohibido gana siempre
    if regla["ambito"] != "PIPELINE":
        p = max(p, 4)
    return p


def gate_de(skill: str, archivo: str, seccion: str, texto: str) -> str | None:
    r = ruta(skill, archivo)
    for patron, cond, gate in GATES:
        if fnmatch.fnmatch(r, patron) and (cond is None or cond.search(seccion) or cond.search(texto)):
            return gate
    return None


# --------------------------------------------------------------------------- build

def insertar_taxonomias(con: sqlite3.Connection, etapas: list[dict]) -> None:
    con.executemany("INSERT OR REPLACE INTO casos (codigo, descripcion, texto_emb) VALUES (?,?,?)",
                    [(k, v, f"caso {k}: {v}") for k, v in CASOS.items()])
    for t in TAREAS:
        con.execute(
            "INSERT OR REPLACE INTO tareas (codigo, etapa, nombre, autoridad, entra, sale, gate, tracks, texto_emb)"
            " VALUES (?,?,?,?,?,?,?,?,?)",
            (t["codigo"], t["etapa"], t["nombre"], t["autoridad"], t["entra"], t["sale"], t["gate"], "",
             f"{t['codigo']} {t['nombre']} · entra: {t['entra']} · sale: {t['sale']}"))
    for e in etapas:
        if e["clave"] in SUELTAS_RESPALDO:
            con.execute(
                "INSERT OR REPLACE INTO tareas (codigo, etapa, nombre, autoridad, entra, sale, gate, tracks, texto_emb)"
                " VALUES (?,?,?,?,?,?,?,?,?)",
                (e["clave"], "SUELTA", e["nombre"], "app/build_app.py", e.get("entra", ""), e.get("sale", ""),
                 e.get("gate", ""), ",".join(e.get("tracks", [])),
                 f"{e['clave']} {e['nombre']} · {e.get('cond', '')} · entra: {e.get('entra', '')} · sale: {e.get('sale', '')}"))
    for dim, meta in FACETAS.items():
        for valor, desc in meta["valores"].items():
            con.execute("INSERT OR REPLACE INTO facetas_catalogo VALUES (?,?,?,?)",
                        (dim, valor, desc, meta["cerrada"]))


def registrar_valor_faceta(con: sqlite3.Connection, dim: str, valor: str) -> None:
    con.execute("INSERT OR IGNORE INTO facetas_catalogo (dimension, valor, descripcion, cerrada) VALUES (?,?,?,?)",
                (dim, valor, f"{FACETAS[dim]['nombre']} {valor}", FACETAS[dim]["cerrada"]))


def cmd_build(args) -> int:
    db = Path(args.db)
    if db.exists():
        if not args.force:
            print(f"{db} ya existe. Usa --force para reconstruirla (se pierden prompts/briefs registrados).",
                  file=sys.stderr)
            return 2
        db.unlink()
    con = connect(str(db))
    con.executescript(SCHEMA.read_text(encoding="utf-8"))
    ba = cargar_build_app()
    insertar_taxonomias(con, ba["ETAPAS"])
    fecha = now()
    extra_roots = [Path(r) for r in args.root]
    skills = list(rr.PIPELINE_SKILLS) + [rr.REGIMENES_SKILL]
    for extra in extra_roots:
        if extra.name not in skills:
            skills.append(extra.name)

    total = 0
    for skill in skills:
        reg = rr.build_v3(skill, extra_roots)
        rules = reg["rules"]
        sha_por_archivo = {a["file"]: a["sha256"] for a in reg["archivos"]}
        con.execute("INSERT OR REPLACE INTO skills VALUES (?,?,?,?,?)",
                    (skill, reg["root"], len(reg["archivos"]), len(rules), fecha))
        for a in reg["archivos"]:
            con.execute("INSERT OR REPLACE INTO archivos VALUES (?,?,?,?,?)",
                        (skill, a["file"], a["sha256"], a["reglas"], fecha))
        for d in reg["descartes"]:
            con.execute("INSERT INTO descartes VALUES (?,?,?,?,?)",
                        (skill, d["file"], d["line"], d["razon"], d["text"]))
        for r in rules:
            gate = gate_de(skill, r["file"], r["seccion"], r["text"])
            texto_emb = f"{skill} · {r['file']} · {r['seccion']} · {r['text']}"
            con.execute(
                "INSERT OR REPLACE INTO reglas (id, skill, archivo, linea, seccion, texto, idioma, marcadores,"
                " fuerza, ambito, medio, fase, prioridad, verificable, gate, estado, fuente, texto_emb,"
                " sha_archivo, extraido) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (r["id"], skill, r["file"], r["line"], r["seccion"], r["text"], r["idioma"],
                 ",".join(r["markers"]), r["fuerza"], r["ambito"],
                 primero(MEDIO, skill, r["file"], "AMBOS"), primero(FASE, skill, r["file"], "TRANSVERSAL"),
                 prioridad_de(r, gate),
                 "MECANICA" if gate else ("NO_VERIFICABLE" if r["ambito"] == "EJEMPLO" else "MANUAL"),
                 gate, r["estado"], r["fuente"], texto_emb, sha_por_archivo.get(r["file"], ""), fecha))
            # facetas heredadas del frontmatter (archivo) o del prefijo de línea (manual)
            for dim, valores in r.get("facetas", {}).items():
                origen = r.get("facetas_origen", {}).get(dim, "archivo")
                for v in valores:
                    registrar_valor_faceta(con, dim, v)
                    con.execute("INSERT OR IGNORE INTO regla_faceta VALUES (?,?,?,?,?)",
                                (r["id"], dim, v, origen, fecha))
            total += 1
        # casos y tareas declarados en el frontmatter de cada archivo de régimen
        if skill not in rr.PIPELINE_SKILLS:
            for root in reg["roots"]:
                for md in rr.archivos_de(Path(root), True):
                    fm = rr.parse_frontmatter(md.read_text(encoding="utf-8").splitlines())
                    rel = str(md.relative_to(root))
                    ids = [r["id"] for r in rules if r["file"] == rel]
                    for c in rr.normalizar_valores(fm.get("casos")):
                        if c in CASOS:
                            for rid in ids:
                                con.execute("INSERT OR IGNORE INTO regla_caso VALUES (?,?,?,?,?,?)",
                                            (rid, c, "archivo", None, "frontmatter", fecha))
                    for t in rr.normalizar_valores(fm.get("tareas")):
                        if con.execute("SELECT 1 FROM tareas WHERE codigo=?", (t,)).fetchone():
                            for rid in ids:
                                con.execute("INSERT OR IGNORE INTO regla_tarea VALUES (?,?,?,?)",
                                            (rid, t, "directa", fecha))
        print(f"  {skill:26} {len(rules):5} reglas · {len(reg['descartes']):3} descartes"
              + ("" if reg["root"] else "  (raíz ausente)"))
    con.commit()
    print(f"reglas: {total} · casos: {len(CASOS)} · tareas: "
          f"{con.execute('SELECT COUNT(*) FROM tareas').fetchone()[0]} · etapas desde {ba['fuente']}")

    if not args.sin_import:
        clas = Path(args.clasificacion)
        if clas.is_file():
            importar_clasificacion(con, clas)
        else:
            print(f"clasificación v2 no encontrada en {clas}; omitida")
    reconstruir_fts(con)
    con.commit()
    return 0


# --------------------------------------------------------------------------- import-app

def importar_clasificacion(con: sqlite3.Connection, path: Path) -> tuple[int, int, int]:
    """Importa .claude/rules/clasificacion_v2_app.json tal cual, conservando el origen.

    Devuelve (reglas en la clasificación, importadas, huérfanas). Una huérfana es un id
    que el registro fresco ya no produce: queda en importacion_huerfana con su texto.
    """
    datos = json.loads(path.read_text(encoding="utf-8"))
    idx: dict = datos.get("idx", {})
    por_caso: dict = datos.get("porCaso", {})
    fecha = now()
    nombre = path.name
    existentes = {r[0] for r in con.execute("SELECT id FROM reglas")}
    importadas: set[str] = set()
    huerfanas: set[str] = set()
    filas = 0
    for caso, ids in por_caso.items():
        for rid in ids:
            meta = idx.get(rid, {})
            origen = meta.get("o", "regla")
            if origen not in ("archivo", "regla", "auditoria"):
                origen = "regla"
            if caso not in CASOS:
                if rid not in huerfanas:
                    con.execute("INSERT INTO importacion_huerfana VALUES (?,?,?,?,?,?,?)",
                                (nombre, rid, caso, origen, "caso_desconocido", json.dumps(meta, ensure_ascii=False), fecha))
                huerfanas.add(rid)
                continue
            if rid not in existentes:
                # un mismo id puede aparecer bajo más de un caso en porCaso; se registra
                # como huérfana una sola vez, no una fila por cada caso donde aparece.
                if rid not in huerfanas:
                    con.execute("INSERT INTO importacion_huerfana VALUES (?,?,?,?,?,?,?)",
                                (nombre, rid, caso, origen, "id_no_en_registro",
                                 f"{meta.get('s')}/{meta.get('a')}:{meta.get('l')} · {meta.get('t', '')[:160]}", fecha))
                huerfanas.add(rid)
                continue
            con.execute("INSERT OR REPLACE INTO regla_caso VALUES (?,?,?,?,?,?)",
                        (rid, caso, origen, None, "app_v2", fecha))
            importadas.add(rid)
            filas += 1
    # ids del índice que no aparecen en ningún caso (no debería pasar; se registra)
    for rid, meta in idx.items():
        if rid not in importadas and rid not in huerfanas:
            con.execute("INSERT INTO importacion_huerfana VALUES (?,?,?,?,?,?,?)",
                        (nombre, rid, None, meta.get("o"), "sin_caso_en_porCaso",
                         json.dumps(meta, ensure_ascii=False), fecha))
            huerfanas.add(rid)
    con.commit()
    total = len(idx)
    print(f"import-app: {total} reglas en {nombre} · {len(importadas)} importadas ({filas} filas regla_caso)"
          f" · {len(huerfanas)} huérfanas (tabla importacion_huerfana)")
    for rid, razon, det in con.execute(
            "SELECT regla_id, razon, detalle FROM importacion_huerfana WHERE origen_import=? GROUP BY regla_id",
            (nombre,)):
        print(f"  huérfana {rid} · {razon} · {det}")
    return total, len(importadas), len(huerfanas)


def cmd_import_app(args) -> int:
    con = connect(args.db)
    total, imp, huer = importar_clasificacion(con, Path(args.clasificacion))
    return 0 if total == imp + huer else 1


# --------------------------------------------------------------------------- derive

def cmd_derive(args) -> int:
    con = connect(args.db)
    fecha = now()
    ba = cargar_build_app()
    caso_etapas = caso_a_etapas(ba["ETAPAS"])
    codigos_por_etapa: dict[str, list[str]] = {}
    for cod, etapa in con.execute("SELECT codigo, etapa FROM tareas"):
        codigos_por_etapa.setdefault(etapa, []).append(cod)
    tareas_validas = {c for c, in con.execute("SELECT codigo FROM tareas")}
    capas = capas_del_director()

    reglas = con.execute("SELECT id, skill, archivo, seccion, texto FROM reglas").fetchall()
    con_caso = {r[0] for r in con.execute("SELECT DISTINCT regla_id FROM regla_caso")}
    n_caso = n_tarea = n_faceta = 0

    def tarea(rid: str, cod: str, origen: str) -> None:
        nonlocal n_tarea
        if cod in tareas_validas:
            cur = con.execute("INSERT OR IGNORE INTO regla_tarea VALUES (?,?,?,?)", (rid, cod, origen, fecha))
            n_tarea += cur.rowcount

    def faceta(rid: str, dim: str, valor: str, origen: str) -> None:
        nonlocal n_faceta
        registrar_valor_faceta(con, dim, valor)
        cur = con.execute("INSERT OR IGNORE INTO regla_faceta VALUES (?,?,?,?,?)", (rid, dim, valor, origen, fecha))
        n_faceta += cur.rowcount

    for rid, skill, archivo, seccion, texto in reglas:
        # 1. caso por ruta, sólo si nadie lo clasificó antes (la auditoría manda)
        if rid not in con_caso:
            r = ruta(skill, archivo)
            for patron, casos in ARCHIVO_CASO:
                if fnmatch.fnmatch(r, patron):
                    for c in casos:
                        cur = con.execute("INSERT OR IGNORE INTO regla_caso VALUES (?,?,?,?,?,?)",
                                          (rid, c, "archivo", None, "derive:ruta", fecha))
                        n_caso += cur.rowcount
                    break
        # 2. tarea directa: códigos citados en el texto
        for cod in set(re.findall(r"\b(E[0-7]\.\d|POST\.[12])\b", texto)):
            tarea(rid, cod, "directa")
        # 3. tarea derivada del caso (mapa caso -> etapa de build_app.py, refinado)
        for (c,) in con.execute("SELECT caso FROM regla_caso WHERE regla_id=?", (rid,)):
            if c == "NINGUNO":
                continue
            refinadas = CASO_SUBPROCESO.get(c)
            for etapa in caso_etapas.get(c, []):
                if etapa in tareas_validas:          # tarea suelta (PROMPT_IMAGEN, …)
                    tarea(rid, etapa, "derivada_caso")
                elif refinadas:
                    for cod in refinadas:
                        tarea(rid, cod, "derivada_caso")
                else:
                    for cod in codigos_por_etapa.get(etapa, []):
                        tarea(rid, cod, "derivada_caso")
        # 4. tarea derivada del archivo (columna "Skill con autoridad")
        for (codigos,) in acumular(ARCHIVO_TAREA, skill, archivo):
            for cod in codigos:
                tarea(rid, cod, "derivada_archivo")
        # 5. tarea derivada de la capa del director §1 (skill -> capa -> etapa completa)
        for n in capas.get(skill, []):
            for cod in codigos_por_etapa.get(f"E{n}", []):
                tarea(rid, cod, "derivada_skill")
        # 6. facetas por ruta y por regex
        for dim, valores in acumular(ARCHIVO_FACETA, skill, archivo):
            for v in valores:
                faceta(rid, dim, v, "archivo")
        for dim, valor, patron in REGEX_FACETA:
            if patron.search(texto):
                faceta(rid, dim, valor, "regex")
    con.commit()
    print(f"derive: +{n_caso} regla_caso (ruta) · +{n_tarea} regla_tarea · +{n_faceta} regla_faceta")
    print(f"  capas del director leídas para {len(capas)} skills: "
          + ", ".join(f"{k}={v}" for k, v in sorted(capas.items())))
    sin = con.execute("SELECT COUNT(*) FROM v_sin_caso_ni_tarea").fetchone()[0]
    print(f"  sin caso ni tarea: {sin}")
    return 0


# --------------------------------------------------------------------------- embed

def reparar_cache_fastembed(cache: Path) -> int:
    """onnxruntime >= 1.20 rechaza que model.onnx_data sea un symlink a otra carpeta
    ("External data path escapes model directory"). fastembed descarga con symlinks a
    blobs compartidos; aquí se reemplazan por hard links (mismo filesystem) o copias."""
    import shutil

    reparados = 0
    for snap in cache.glob("models--*/snapshots/*"):
        for f in snap.iterdir():
            if f.is_symlink():
                destino = f.resolve()
                f.unlink()
                try:
                    os.link(destino, f)
                except OSError:
                    shutil.copy2(destino, f)
                reparados += 1
    return reparados


def cargar_modelo():
    try:
        from fastembed import TextEmbedding  # type: ignore
    except ImportError:
        print("embed: falta fastembed (usa el venv). pip install fastembed numpy", file=sys.stderr)
        return None
    FASTEMBED_CACHE.mkdir(parents=True, exist_ok=True)
    try:
        return TextEmbedding(MODELO_EMB, cache_dir=str(FASTEMBED_CACHE))
    except Exception as exc:  # noqa: BLE001
        if "External data path" in str(exc):
            n = reparar_cache_fastembed(FASTEMBED_CACHE)
            print(f"embed: caché reparado ({n} symlinks -> hard links); reintentando")
            return TextEmbedding(MODELO_EMB, cache_dir=str(FASTEMBED_CACHE))
        raise


def vectorizar(modelo, textos: list[str], lote: int = 32):
    import numpy as np

    vs = np.asarray(list(modelo.embed(textos, batch_size=lote)), dtype=np.float32)
    normas = np.linalg.norm(vs, axis=1, keepdims=True)
    normas[normas == 0] = 1.0
    return vs / normas


def cmd_embed(args) -> int:
    import numpy as np  # noqa: F401

    con = connect(args.db)
    t0 = time.time()
    modelo = cargar_modelo()
    if modelo is None:
        return 2
    t_carga = time.time() - t0

    # re-embed sólo lo que no tiene vector o cuyo texto cambió (texto_sha distinto)
    hechos = dict(con.execute("SELECT regla_id, texto_sha FROM embeddings WHERE modelo_emb=?", (MODELO_EMB,)))
    pendientes = [(rid, t) for rid, t in con.execute("SELECT id, texto_emb FROM reglas")
                  if hechos.get(rid) != sha1(t)]
    t1 = time.time()
    hechas = 0
    for i in range(0, len(pendientes), args.lote):
        bloque = pendientes[i:i + args.lote]
        vs = vectorizar(modelo, ["passage: " + t for _, t in bloque], args.lote)
        for (rid, t), v in zip(bloque, vs):
            con.execute("INSERT OR REPLACE INTO embeddings VALUES (?,?,?,?,?)",
                        (rid, MODELO_EMB, int(v.shape[0]), v.astype("<f4").tobytes(), sha1(t)))
        hechas += len(bloque)
        con.commit()
    t_reglas = time.time() - t1

    # prototipos: casos, tareas y valores de faceta (centroide = una descripción cada uno)
    protos: list[tuple[str, str, str]] = []
    for cod, desc, emb in con.execute("SELECT codigo, descripcion, texto_emb FROM casos"):
        protos.append(("caso", cod, emb or f"caso {cod}: {desc}"))
    for cod, emb in con.execute("SELECT codigo, texto_emb FROM tareas"):
        protos.append(("tarea", cod, emb or cod))
    for dim, valor, desc in con.execute("SELECT dimension, valor, descripcion FROM facetas_catalogo WHERE cerrada=1"):
        protos.append(("faceta", f"{dim}:{valor}", f"{FACETAS[dim]['nombre']} {valor}: {desc}"))
    t2 = time.time()
    vs = vectorizar(modelo, ["passage: " + t for _, _, t in protos], args.lote)
    for (tipo, cod, t), v in zip(protos, vs):
        con.execute("INSERT OR REPLACE INTO prototipos VALUES (?,?,?,?,?,?)",
                    (tipo, cod, MODELO_EMB, int(v.shape[0]), v.astype("<f4").tobytes(), t))
    con.commit()
    total = con.execute("SELECT COUNT(*) FROM embeddings WHERE modelo_emb=?", (MODELO_EMB,)).fetchone()[0]
    print(f"embed: modelo {MODELO_EMB} cargado en {t_carga:.1f}s · {hechas} reglas vectorizadas en "
          f"{t_reglas:.1f}s · {len(protos)} prototipos en {time.time() - t2:.1f}s · "
          f"{total} embeddings en la base · total {time.time() - t0:.1f}s")
    return 0


# --------------------------------------------------------------------------- fts

def reconstruir_fts(con: sqlite3.Connection) -> int:
    con.execute("DELETE FROM reglas_fts")
    filas = con.execute("SELECT id, skill, seccion, texto FROM reglas").fetchall()
    con.executemany("INSERT INTO reglas_fts (id, skill, seccion, texto) VALUES (?,?,?,?)", filas)
    con.commit()
    print(f"fts: {len(filas)} reglas indexadas (unicode61 remove_diacritics 2)")
    return len(filas)


def cmd_fts(args) -> int:
    reconstruir_fts(connect(args.db))
    return 0


# --------------------------------------------------------------------------- classify (capa 2: vector)
# Decisión 9 (2026-09-25): NotebookLM queda fuera del flujo por completo. Quién decide
# qué reglas aplican a una faceta o caso se reparte en tres capas, en este orden de
# autoridad (la de abajo nunca puede ser pisada por las de arriba):
#   1. Determinista — ya la aplica `derive`: ARCHIVO_CASO/ARCHIVO_FACETA/REGEX_FACETA en
#      este mismo archivo, con la disciplina de scope.yaml (primer patrón que hace match
#      manda, cada entrada se puede discutir línea por línea). origen='archivo'|'regex'.
#   2. Vectorial — este comando. Solo toca lo que la capa 1 dejó sin asignar. Compara el
#      embedding de la regla contra el prototipo de cada valor posible (construido por
#      `embed` a partir de las descripciones de FACETAS/CASOS) y asigna origen='vector'
#      con la confianza (coseno) si supera --umbral (0.80 por defecto). Es una sugerencia
#      con número, nunca una decisión editorial.
#   3. Auditoría de Eric — `audit-export`/`audit-import`. Todo lo que sigue sin valor
#      tras las capas 1 y 2 sale a una hoja xlsx; su corrección entra con origen='auditoria'
#      y manda sobre las otras dos, sin importar la confianza que tuvieran.
#
# Las facetas D1-D9 de la tarjeta de ancla solo tienen sentido para reglas de imagen o
# video (Decisión 5); el resto del corpus (guion, empaquetado, meta) no se audita en D1-D9.
FACETA_MEDIOS = ("IMAGEN", "VIDEO", "AMBOS")
DIMENSIONES_CERRADAS = ("d1", "d2", "d3", "d4", "d8", "d9")


def _vec(blob: bytes):
    import numpy as np

    return np.frombuffer(blob, dtype="<f4")


def _mejor(v, candidatos: list[tuple[str, "object"]]) -> tuple[str | None, float]:
    import numpy as np

    if not candidatos:
        return None, -1.0
    mejor_valor, mejor_score = None, -1.0
    for valor, cv in candidatos:
        score = float(np.dot(v, cv))
        if score > mejor_score:
            mejor_valor, mejor_score = valor, score
    return mejor_valor, mejor_score


def _cargar_prototipos(con: sqlite3.Connection):
    protos_caso: list[tuple[str, object]] = []
    protos_faceta: dict[str, list[tuple[str, object]]] = {}
    for tipo, codigo, blob in con.execute(
            "SELECT tipo, codigo, vector FROM prototipos WHERE modelo_emb=?", (MODELO_EMB,)):
        v = _vec(blob)
        if tipo == "caso":
            protos_caso.append((codigo, v))
        elif tipo == "faceta":
            dim, valor = codigo.split(":", 1)
            protos_faceta.setdefault(dim, []).append((valor, v))
    return protos_caso, protos_faceta


def cmd_classify(args) -> int:
    con = connect(args.db)
    fecha = now()
    umbral = args.umbral

    vec_by_id = {rid: _vec(blob) for rid, blob in
                 con.execute("SELECT regla_id, vector FROM embeddings WHERE modelo_emb=?", (MODELO_EMB,))}
    if not vec_by_id:
        print("classify: no hay embeddings en la base. Corre `embed` antes de `classify` "
              "(capa 1 determinista ya corrió con `derive`; sin vectores, capa 2 no puede sugerir "
              "nada y todo lo pendiente va directo a auditoría).", file=sys.stderr)
        return 2

    protos_caso, protos_faceta = _cargar_prototipos(con)
    con_caso = {r[0] for r in con.execute("SELECT DISTINCT regla_id FROM regla_caso")}
    medio_por_regla = dict(con.execute("SELECT id, medio FROM reglas"))
    facetas_por_regla: dict[str, set[str]] = {}
    for rid, dim in con.execute("SELECT regla_id, dimension FROM regla_faceta"):
        facetas_por_regla.setdefault(rid, set()).add(dim)

    n_caso = n_caso_bajo_umbral = 0
    n_faceta = n_faceta_bajo_umbral = 0
    for rid, v in vec_by_id.items():
        if rid not in con_caso:
            cod, score = _mejor(v, protos_caso)
            if cod is not None and score >= umbral:
                con.execute("INSERT OR IGNORE INTO regla_caso VALUES (?,?,?,?,?,?)",
                            (rid, cod, "vector", score, None, fecha))
                n_caso += 1
            else:
                n_caso_bajo_umbral += 1
        if medio_por_regla.get(rid) in FACETA_MEDIOS:
            tiene = facetas_por_regla.get(rid, set())
            for dim in DIMENSIONES_CERRADAS:
                if dim in tiene:
                    continue
                valor, score = _mejor(v, protos_faceta.get(dim, []))
                if valor is not None and score >= umbral:
                    registrar_valor_faceta(con, dim, valor)
                    con.execute("INSERT OR IGNORE INTO regla_faceta VALUES (?,?,?,?,?)",
                                (rid, dim, valor, "vector", fecha))
                    n_faceta += 1
                else:
                    n_faceta_bajo_umbral += 1
    con.commit()
    print(f"classify: capa vectorial, umbral {umbral}")
    print(f"  regla_caso   +{n_caso} asignadas por vector · {n_caso_bajo_umbral} bajo umbral o sin candidato (van a auditoría)")
    print(f"  regla_faceta +{n_faceta} asignadas por vector · {n_faceta_bajo_umbral} bajo umbral o sin candidato (van a auditoría)")
    return 0


# --------------------------------------------------------------------------- auditoría (capa 3: Eric, xlsx)

def _texto_regla(con: sqlite3.Connection, rid: str) -> tuple[str, str, str, int, str]:
    row = con.execute("SELECT skill, archivo, linea, seccion, texto FROM reglas WHERE id=?", (rid,)).fetchone()
    return row if row else ("", "", 0, "", "")


def cmd_audit_export(args) -> int:
    try:
        from openpyxl import Workbook
    except ImportError:
        print("audit-export: falta openpyxl. pip install openpyxl", file=sys.stderr)
        return 2

    con = connect(args.db)
    umbral = args.umbral
    protos_caso, protos_faceta = _cargar_prototipos(con)
    vec_by_id = {rid: _vec(blob) for rid, blob in
                 con.execute("SELECT regla_id, vector FROM embeddings WHERE modelo_emb=?", (MODELO_EMB,))}
    medio_por_regla = dict(con.execute("SELECT id, medio FROM reglas"))
    con_caso = {r[0] for r in con.execute("SELECT DISTINCT regla_id FROM regla_caso")}
    facetas_por_regla: dict[str, set[str]] = {}
    for rid, dim in con.execute("SELECT regla_id, dimension FROM regla_faceta"):
        facetas_por_regla.setdefault(rid, set()).add(dim)

    wb = Workbook()
    wb.remove(wb.active)
    cabecera = ["id", "skill", "archivo", "linea", "seccion", "texto", "valor_sugerido", "confianza",
                "valor_corregido", "razon"]

    def hoja(nombre: str) -> None:
        ws = wb.create_sheet(nombre[:31])
        ws.append(cabecera)

    # hoja 'caso': toda regla sin caso tras las capas 1 y 2
    pendientes_caso = [rid for (rid,) in con.execute("SELECT id FROM reglas") if rid not in con_caso]
    hoja("caso")
    ws = wb["caso"]
    for rid in pendientes_caso:
        skill, archivo, linea, seccion, texto = _texto_regla(con, rid)
        v = vec_by_id.get(rid)
        sugerido, score = (_mejor(v, protos_caso) if v is not None else (None, None))
        ws.append([rid, skill, archivo, linea, seccion, texto[:500], sugerido or "",
                   round(score, 4) if score is not None else "", "", ""])

    # una hoja por faceta cerrada, solo reglas de imagen/video
    for dim in DIMENSIONES_CERRADAS:
        hoja(f"faceta_{dim}")
        ws = wb[f"faceta_{dim}"]
        candidatos_dim = protos_faceta.get(dim, [])
        for (rid,) in con.execute("SELECT id FROM reglas"):
            if medio_por_regla.get(rid) not in FACETA_MEDIOS:
                continue
            if dim in facetas_por_regla.get(rid, set()):
                continue
            skill, archivo, linea, seccion, texto = _texto_regla(con, rid)
            v = vec_by_id.get(rid)
            sugerido, score = (_mejor(v, candidatos_dim) if v is not None else (None, None))
            ws.append([rid, skill, archivo, linea, seccion, texto[:500], sugerido or "",
                       round(score, 4) if score is not None else "", "", ""])

    wb.save(args.xlsx)
    resumen = ", ".join(f"{s.title}={s.max_row - 1}" for s in wb.worksheets)
    print(f"audit-export: {args.xlsx} escrito · filas por hoja: {resumen}")
    print(f"  umbral de referencia para lo ya auto-asignado por vector: {umbral} "
          "(esta hoja es justo lo que quedó por debajo o sin candidato)")
    return 0


def cmd_audit_import(args) -> int:
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("audit-import: falta openpyxl. pip install openpyxl", file=sys.stderr)
        return 2

    con = connect(args.db)
    fecha = now()
    auditor = args.auditor
    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    n_caso = n_faceta = 0

    for ws in wb.worksheets:
        cab = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {name: i for i, name in enumerate(cab)}
        requeridos = {"id", "valor_sugerido", "valor_corregido", "razon"}
        if not requeridos.issubset(idx):
            print(f"audit-import: hoja '{ws.title}' sin las columnas esperadas, se omite", file=sys.stderr)
            continue
        for fila in ws.iter_rows(min_row=2, values_only=True):
            rid = fila[idx["id"]]
            corregido = fila[idx["valor_corregido"]]
            if not rid or not corregido:
                continue
            corregido = str(corregido).strip()
            if not corregido:
                continue
            sugerido = fila[idx["valor_sugerido"]] or ""
            razon = fila[idx["razon"]] or ""
            if ws.title == "caso":
                con.execute("INSERT OR REPLACE INTO regla_caso VALUES (?,?,?,?,?,?)",
                            (rid, corregido, "auditoria", 1.0, auditor, fecha))
                con.execute("INSERT INTO auditorias VALUES (?,?,?,?,?,?)",
                            (rid, str(sugerido), corregido, razon, auditor, fecha))
                n_caso += 1
            elif ws.title.startswith("faceta_"):
                dim = ws.title.removeprefix("faceta_")
                registrar_valor_faceta(con, dim, corregido)
                con.execute("INSERT OR REPLACE INTO regla_faceta VALUES (?,?,?,?,?)",
                            (rid, dim, corregido, "auditoria", fecha))
                con.execute("INSERT INTO auditorias VALUES (?,?,?,?,?,?)",
                            (rid, f"{dim}:{sugerido}", f"{dim}:{corregido}", razon, auditor, fecha))
                n_faceta += 1
    con.commit()
    print(f"audit-import: {args.xlsx} · +{n_caso} regla_caso corregidas · +{n_faceta} regla_faceta corregidas "
          f"(origen='auditoria', auditor={auditor})")
    return 0


# --------------------------------------------------------------------------- stats / check

def cmd_stats(args) -> int:
    con = connect(args.db)
    q = lambda sql, *p: con.execute(sql, p).fetchone()[0]  # noqa: E731
    tot = q("SELECT COUNT(*) FROM reglas")
    print(f"reglas            {tot}")
    print(f"con caso          {q('SELECT COUNT(DISTINCT regla_id) FROM regla_caso')}")
    print(f"con tarea         {q('SELECT COUNT(DISTINCT regla_id) FROM regla_tarea')}")
    print(f"con faceta d1-d9  {q('SELECT COUNT(DISTINCT regla_id) FROM regla_faceta')}")
    print(f"con embedding     {q('SELECT COUNT(DISTINCT regla_id) FROM embeddings')}")
    print(f"en v_sin_faceta   {q('SELECT COUNT(*) FROM v_sin_faceta')}   (falta caso, tarea o facetas)")
    print(f"sin caso ni tarea {q('SELECT COUNT(*) FROM v_sin_caso_ni_tarea')}")
    print(f"descartes         {q('SELECT COUNT(*) FROM descartes')}   huérfanas de importación: "
          f"{q('SELECT COUNT(DISTINCT regla_id) FROM importacion_huerfana')}")
    for campo in ("estado", "fuerza", "ambito", "idioma", "medio", "fase", "prioridad"):
        print(f"\npor {campo}:")
        for v, n in con.execute(f"SELECT {campo}, COUNT(*) FROM reglas GROUP BY {campo} ORDER BY 2 DESC"):
            print(f"  {str(v):16} {n:>5}")
    print("\npor origen de clasificación (caso):")
    for o, n in con.execute("SELECT origen, COUNT(DISTINCT regla_id) FROM regla_caso GROUP BY origen ORDER BY 2 DESC"):
        print(f"  {o:16} {n:>5}")
    print("\nreglas por caso:")
    for cod, desc, n in con.execute("SELECT * FROM v_por_caso"):
        print(f"  {cod:9} {n:>5}   {desc[:60]}")
    print("\nreglas por tarea:")
    for cod, nombre, n in con.execute("SELECT * FROM v_por_tarea"):
        print(f"  {cod:13} {n:>5}   {nombre[:50]}")
    print("\nfacetas por dimensión (reglas distintas):")
    for dim, n, vals in con.execute(
            "SELECT dimension, COUNT(DISTINCT regla_id), COUNT(DISTINCT valor) FROM regla_faceta GROUP BY dimension"):
        print(f"  {dim} {FACETAS[dim]['nombre']:15} {n:>5} reglas · {vals} valores")
    print("\npor skill:")
    for skill, n in con.execute("SELECT skill, COUNT(*) FROM reglas GROUP BY skill ORDER BY 2 DESC"):
        print(f"  {skill:26} {n:>5}")
    print(f"\nbriefs registrados: {q('SELECT COUNT(*) FROM briefs')} · auditorías: {q('SELECT COUNT(*) FROM auditorias')}")
    return 0


def cmd_check(args) -> int:
    con = connect(args.db)
    print("v_sin_faceta por skill y por lo que falta:")
    for skill, falta, n in con.execute(
            "SELECT skill, falta, COUNT(*) FROM v_sin_faceta GROUP BY skill, falta ORDER BY skill, falta"):
        print(f"  {skill:26} {falta:22} {n:>5}")
    total_sf = con.execute("SELECT COUNT(*) FROM v_sin_faceta").fetchone()[0]
    print(f"  total en v_sin_faceta: {total_sf}")
    huer = con.execute("SELECT COUNT(DISTINCT regla_id) FROM importacion_huerfana").fetchone()[0]
    if huer:
        print(f"huérfanas de importación: {huer} (visibles en importacion_huerfana)")
    filas = con.execute("SELECT id, skill, archivo, linea, texto FROM v_sin_caso_ni_tarea").fetchall()
    if filas:
        print(f"\nFALLA: {len(filas)} reglas sin caso NI tarea — no hay forma de servirlas:")
        for rid, skill, archivo, linea, texto in filas[:50]:
            print(f"  {rid} {skill}/{archivo}:{linea} · {texto[:90]}")
        return 1
    print("\nOK: 0 reglas sin caso ni tarea. Toda regla es alcanzable por al menos una dimensión.")
    return 0


# --------------------------------------------------------------------------- cli

def main() -> int:
    ap = argparse.ArgumentParser(description="Base de reglas v3 del pipeline visual")
    ap.add_argument("--db", default="rules.sqlite")
    sub = ap.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build", help="esquema + registro + import-app + fts (base nueva)")
    b.add_argument("--force", action="store_true", help="borra la base si existe")
    b.add_argument("--root", action="append", default=[], help="raíz extra de régimen (repetible)")
    b.add_argument("--clasificacion", default=str(CLASIFICACION_V2))
    b.add_argument("--sin-import", action="store_true")
    b.set_defaults(fn=cmd_build)

    i = sub.add_parser("import-app", help="importa la clasificación v2 de la app")
    i.add_argument("--clasificacion", default=str(CLASIFICACION_V2))
    i.set_defaults(fn=cmd_import_app)

    d = sub.add_parser("derive", help="tareas, casos por ruta y facetas")
    d.set_defaults(fn=cmd_derive)

    e = sub.add_parser("embed", help="vectores con fastembed (requiere venv)")
    e.add_argument("--lote", type=int, default=32)
    e.set_defaults(fn=cmd_embed)

    f = sub.add_parser("fts", help="reconstruye reglas_fts")
    f.set_defaults(fn=cmd_fts)

    cl = sub.add_parser("classify", help="capa 2: vector contra prototipos, solo lo que capa 1 dejó sin asignar")
    cl.add_argument("--umbral", type=float, default=0.80)
    cl.set_defaults(fn=cmd_classify)

    ax = sub.add_parser("audit-export", help="capa 3: vuelca a xlsx lo que quedó sin caso/faceta")
    ax.add_argument("--xlsx", default="auditoria.xlsx")
    ax.add_argument("--umbral", type=float, default=0.80)
    ax.set_defaults(fn=cmd_audit_export)

    ai = sub.add_parser("audit-import", help="capa 3: aplica las correcciones de Eric, origen='auditoria'")
    ai.add_argument("--xlsx", default="auditoria.xlsx")
    ai.add_argument("--auditor", default="eric")
    ai.set_defaults(fn=cmd_audit_import)

    s = sub.add_parser("stats")
    s.set_defaults(fn=cmd_stats)

    c = sub.add_parser("check", help="exit 1 si hay reglas sin caso ni tarea")
    c.set_defaults(fn=cmd_check)

    args = ap.parse_args()
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(main())
